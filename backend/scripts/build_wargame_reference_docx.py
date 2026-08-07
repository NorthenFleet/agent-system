from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "assets" / "wargame_a4_reference.docx"
BODY_FONT = "Songti SC"
HEADING_FONT = "Heiti SC"


def set_east_asia_font(style, font_name: str) -> None:
    style.font.name = font_name
    fonts = style.element.rPr.rFonts
    fonts.set(qn("w:ascii"), font_name)
    fonts.set(qn("w:hAnsi"), font_name)
    fonts.set(qn("w:eastAsia"), font_name)
    fonts.set(qn("w:cs"), font_name)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)


def rewrite_theme_fonts(path: Path, font_name: str) -> None:
    with NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(path, "r") as source, ZipFile(
            temporary,
            "w",
            compression=ZIP_DEFLATED,
        ) as target:
            for entry in source.infolist():
                payload = source.read(entry.filename)
                if entry.filename == "word/theme/theme1.xml":
                    xml = payload.decode("utf-8")
                    xml = xml.replace(
                        '<a:ea typeface=""/>',
                        f'<a:ea typeface="{font_name}"/>',
                    )
                    xml = xml.replace(
                        '<a:font script="Hans" typeface="宋体"/>',
                        f'<a:font script="Hans" typeface="{font_name}"/>',
                    )
                    payload = xml.encode("utf-8")
                target.writestr(entry, payload)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


document = Document()
section = document.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.orientation = WD_ORIENT.PORTRAIT
section.left_margin = Mm(25)
section.right_margin = Mm(20)
section.top_margin = Mm(20)
section.bottom_margin = Mm(20)
section.header_distance = Mm(10)
section.footer_distance = Mm(10)

normal = document.styles["Normal"]
set_east_asia_font(normal, BODY_FONT)
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.first_line_indent = Pt(21)
normal.paragraph_format.widow_control = True

title = document.styles["Title"]
set_east_asia_font(title, HEADING_FONT)
title.font.size = Pt(22)
title.font.bold = True
title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(80)
title.paragraph_format.space_after = Pt(24)

subtitle = document.styles["Subtitle"]
set_east_asia_font(subtitle, HEADING_FONT)
subtitle.font.size = Pt(14)
subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)

for name, size, centered, page_break in [
    ("Heading 1", 18, True, True),
    ("Heading 2", 15, False, False),
    ("Heading 3", 13, False, False),
    ("Heading 4", 11, False, False),
]:
    style = document.styles[name]
    set_east_asia_font(style, HEADING_FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    style.paragraph_format.space_before = Pt(14)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.page_break_before = page_break

for name in ["TOC 1", "TOC 2", "TOC 3"]:
    try:
        style = document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    set_east_asia_font(style, BODY_FONT)
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(3)

if "Caption" in document.styles:
    caption = document.styles["Caption"]
    set_east_asia_font(caption, BODY_FONT)
    caption.font.size = Pt(9)
    caption.font.bold = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

table_style = document.styles["Table Grid"]
set_east_asia_font(table_style, BODY_FONT)
table_style.font.size = Pt(9)

header = section.header
header_paragraph = header.paragraphs[0]
header_paragraph.text = "谋战·水面舰艇编队战术手工兵棋"
header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_paragraph.style = normal
header_paragraph.runs[0].font.size = Pt(9)

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_paragraph.add_run("第 ")
add_field(footer_paragraph, "PAGE")
footer_paragraph.add_run(" 页，共 ")
add_field(footer_paragraph, "NUMPAGES")
footer_paragraph.add_run(" 页")
for run in footer_paragraph.runs:
    run.font.size = Pt(9)

document.add_paragraph("参考模板", style="Title")
document.add_paragraph("该内容在 Pandoc 使用参考模板时不会进入正式文档。", style="Subtitle")
document.add_heading("一级标题", level=1)
document.add_heading("二级标题", level=2)
document.add_heading("三级标题", level=3)
document.add_heading("四级标题", level=4)
document.add_paragraph("正文样式示例。")
table = document.add_table(rows=2, cols=3)
table.style = "Table Grid"
for index, text in enumerate(["表头一", "表头二", "表头三"]):
    table.rows[0].cells[index].text = text
set_repeat_table_header(table.rows[0])
for index, text in enumerate(["数据一", "数据二", "数据三"]):
    table.rows[1].cells[index].text = text

core = document.core_properties
core.title = "谋战·水面舰艇编队战术手工兵棋A4排版模板"
core.subject = "规则手册、裁决规则、算子表Word正式输出"
core.author = "OpenClaw文档撰写模块"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
rewrite_theme_fonts(OUTPUT, BODY_FONT)
print(OUTPUT)
