"""Pluggable storage, scanning, OCR, verification and statement adapters."""
from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Protocol


ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "image/jpeg",
    "image/png",
    "image/tiff",
}
MAX_ATTACHMENT_BYTES = int(os.getenv("FINANCE_ATTACHMENT_MAX_BYTES", str(20 * 1024 * 1024)))


class ObjectStorage(Protocol):
    def put(self, object_key: str, content: bytes, content_type: str) -> None: ...
    def read(self, object_key: str) -> bytes: ...
    def signed_url(self, object_key: str, expires_seconds: int = 300) -> str: ...
    def health(self) -> None: ...


class LocalObjectStorage:
    def __init__(self) -> None:
        self.root = Path(os.getenv("FINANCE_OBJECT_ROOT", "data/finance-objects")).resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def put(self, object_key: str, content: bytes, content_type: str) -> None:
        target = (self.root / object_key).resolve()
        if self.root not in target.parents:
            raise ValueError("非法对象路径")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(content)

    def signed_url(self, object_key: str, expires_seconds: int = 300) -> str:
        # Local downloads are authenticated by the finance API rather than by a
        # public filesystem mount.
        return f"/api/finance/attachments/download/{object_key}"

    def read(self, object_key: str) -> bytes:
        target = (self.root / object_key).resolve()
        if self.root not in target.parents or not target.is_file():
            raise FileNotFoundError(object_key)
        return target.read_bytes()

    def health(self) -> None:
        if not self.root.is_dir() or not os.access(self.root, os.R_OK | os.W_OK | os.X_OK):
            raise RuntimeError("本地对象存储目录不可读写")


class S3ObjectStorage:
    def __init__(self) -> None:
        try:
            import boto3
        except ImportError as exc:  # pragma: no cover - deployment guard
            raise RuntimeError("FINANCE_STORAGE=s3 需要安装 boto3") from exc
        self.bucket = os.environ["FINANCE_S3_BUCKET"]
        self.client = boto3.client(
            "s3",
            endpoint_url=os.getenv("FINANCE_S3_ENDPOINT"),
            region_name=os.getenv("FINANCE_S3_REGION", "us-east-1"),
        )

    def put(self, object_key: str, content: bytes, content_type: str) -> None:
        self.client.put_object(Bucket=self.bucket, Key=object_key, Body=content, ContentType=content_type)

    def signed_url(self, object_key: str, expires_seconds: int = 300) -> str:
        return self.client.generate_presigned_url(
            "get_object",
            Params={"Bucket": self.bucket, "Key": object_key},
            ExpiresIn=expires_seconds,
        )

    def read(self, object_key: str) -> bytes:
        response = self.client.get_object(Bucket=self.bucket, Key=object_key)
        return response["Body"].read()

    def health(self) -> None:
        self.client.head_bucket(Bucket=self.bucket)


def object_storage() -> ObjectStorage:
    return S3ObjectStorage() if os.getenv("FINANCE_STORAGE", "local").lower() == "s3" else LocalObjectStorage()


class AttachmentScanner:
    def validate(self, filename: str, content_type: str, content: bytes) -> str:
        if content_type not in ALLOWED_CONTENT_TYPES:
            raise ValueError(f"不支持的附件类型: {content_type}")
        if not content or len(content) > MAX_ATTACHMENT_BYTES:
            raise ValueError("附件为空或超过大小限制")
        lower = filename.lower()
        if lower.endswith((".exe", ".js", ".html", ".svg", ".sh")):
            raise ValueError("附件扩展名不安全")
        # Deployments can replace this adapter with ClamAV.  The built-in
        # scanner intentionally fails closed on obvious executable signatures.
        if content.startswith((b"MZ", b"\x7fELF")):
            raise ValueError("附件包含可执行文件签名")
        return "clean"


class InvoiceOCRAdapter(Protocol):
    def extract(self, content: bytes, content_type: str) -> dict[str, Any]: ...


class ManualReviewOCR:
    def extract(self, content: bytes, content_type: str) -> dict[str, Any]:
        return {"status": "manual_review", "fields": {}, "reason": "OCR provider not configured"}


class InvoiceVerificationAdapter(Protocol):
    def verify(self, invoice: dict[str, Any]) -> dict[str, Any]: ...


class ManualVerification:
    def verify(self, invoice: dict[str, Any]) -> dict[str, Any]:
        return {"status": "manual_review", "reason": "verification provider not configured"}


def ocr_adapter() -> InvoiceOCRAdapter:
    return ManualReviewOCR()


def verification_adapter() -> InvoiceVerificationAdapter:
    return ManualVerification()


@dataclass(frozen=True)
class StatementRow:
    transaction_ref: str
    transaction_date: date
    amount: Decimal
    counterparty: str
    account_masked: str
    memo: str
    raw_payload: dict[str, Any]


def _parse_date(value: str) -> date:
    raw = value.strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"无法识别交易日期: {raw}")


def _parse_amount(value: str) -> Decimal:
    try:
        amount = Decimal(value.replace(",", "").strip()).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise ValueError(f"无法识别交易金额: {value}") from exc
    if amount == 0:
        raise ValueError("交易金额不能为 0")
    return amount


def parse_statement(filename: str, content: bytes) -> list[StatementRow]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig")
        records = list(csv.DictReader(io.StringIO(text)))
    elif lower.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - deployment guard
            raise ValueError("XLSX 流水导入需要安装 openpyxl") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        records = [dict(zip(headers, row)) for row in rows[1:]]
    else:
        raise ValueError("银行流水仅支持 CSV 或 XLSX")

    aliases = {
        "transaction_ref": ("transaction_ref", "流水号", "交易流水号"),
        "transaction_date": ("transaction_date", "交易日期", "日期"),
        "amount": ("amount", "金额", "交易金额"),
        "counterparty": ("counterparty", "对方户名", "收款方"),
        "account": ("account", "对方账号", "账号"),
        "memo": ("memo", "摘要", "备注"),
    }

    def pick(row: dict[str, Any], key: str) -> str:
        for candidate in aliases[key]:
            if candidate in row and row[candidate] not in (None, ""):
                return str(row[candidate])
        return ""

    parsed: list[StatementRow] = []
    for index, row in enumerate(records, start=2):
        ref = pick(row, "transaction_ref") or f"ROW-{index}"
        account = pick(row, "account")
        safe_payload = {}
        for key, value in row.items():
            name = str(key)
            text_value = str(value or "")
            if any(marker in name.lower() for marker in ("account", "账号", "卡号")):
                digits = "".join(char for char in text_value if char.isdigit())
                text_value = ("*" * max(0, len(digits) - 4) + digits[-4:]) if digits else ""
            safe_payload[name] = text_value
        parsed.append(StatementRow(
            transaction_ref=ref,
            transaction_date=_parse_date(pick(row, "transaction_date")),
            amount=_parse_amount(pick(row, "amount")),
            counterparty=pick(row, "counterparty"),
            account_masked=("*" * max(0, len(account) - 4) + account[-4:]) if account else "",
            memo=pick(row, "memo"),
            raw_payload=safe_payload,
        ))
    return parsed
