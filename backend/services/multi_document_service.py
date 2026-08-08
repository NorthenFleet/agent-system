"""Project-level multi-document workspaces for rich text, workbooks and presentations."""

from __future__ import annotations

import copy
import fcntl
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from services.document_workspace_service import (
    DocumentVersionConflict,
    DocumentWorkspaceError,
    document_workspace_service,
)


DOCUMENT_KINDS = {"rich_text", "workbook", "presentation"}
OUTPUT_FORMATS = {"docx", "pdf", "pptx"}
PUBLICATION_STATUSES = {"draft", "review", "approved", "published", "internal"}
STRUCTURE_BINDING_MODES = {"canonical", "derived", "mapped"}
STRUCTURE_BINDING_STATUSES = {"aligned", "diverged", "stale", "missing"}
EDIT_POLICIES = {"editable", "read_only"}
DELIVERY_ROLES = {"deliverable", "historical_reference"}
LINEAGE_SOURCE_TYPES = {"markdown", "chapter_bundle", "structured_authority"}
MAX_UPLOAD_BYTES = 100 * 1024 * 1024
MAX_ASSET_UPLOAD_BYTES = 20 * 1024 * 1024
IMAGE_ASSET_TYPES = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/webp": ".webp",
    "image/gif": ".gif",
    "image/svg+xml": ".svg",
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_name(value: str) -> str:
    cleaned = "".join(
        ch if ch.isalnum() or ch in {"-", "_", " ", "·", "—"} else "-"
        for ch in value.strip()
    )
    return cleaned.strip(" -") or "document"


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _docx_heading_level(style_name: str) -> int:
    value = str(style_name or "").strip().lower()
    match = re.search(r"(?:heading|标题)\s*([1-6])", value)
    return int(match.group(1)) if match else 0


def _docx_table_markdown(table: Any) -> str:
    rows = []
    for row in table.rows:
        values = []
        for cell in row.cells:
            value = " ".join(cell.text.split()).replace("|", "\\|")
            values.append(value)
        rows.append(values)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = "| " + " | ".join(rows[0]) + " |"
    divider = "| " + " | ".join(["---"] * width) + " |"
    body = ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return "\n".join([header, divider, *body])


def _convert_docx(data: bytes, title: str) -> dict[str, Any]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn
    except Exception as exc:  # pragma: no cover - deployment dependency guard
        raise DocumentWorkspaceError("服务器缺少 python-docx，无法导入 Word") from exc

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentWorkspaceError("Word 文件解析失败") from exc

    lines = ["---", f"title: {json.dumps(title, ensure_ascii=False)}", "status: imported", "---", ""]
    assets: dict[str, bytes] = {}
    relationship_names: dict[str, str] = {}
    paragraph_count = 0
    heading_count = 0
    table_count = 0
    list_count = 0
    image_count = 0

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = " ".join(paragraph.text.split())
            image_lines = []
            for blip in paragraph._p.xpath(".//a:blip"):
                relationship_id = str(blip.get(qn("r:embed")) or "")
                if not relationship_id:
                    continue
                if relationship_id not in relationship_names:
                    part = document.part.related_parts.get(relationship_id)
                    if part is None or not hasattr(part, "blob"):
                        continue
                    suffix = Path(str(getattr(part, "partname", ""))).suffix.lower() or ".bin"
                    filename = f"image-{len(relationship_names) + 1:03d}{suffix}"
                    relationship_names[relationship_id] = filename
                    assets[filename] = bytes(part.blob)
                filename = relationship_names[relationship_id]
                image_count += 1
                image_lines.append(f"![图片{image_count}](assets/{filename})")
            if not text and not image_lines:
                continue
            level = _docx_heading_level(getattr(paragraph.style, "name", ""))
            style_name = str(getattr(paragraph.style, "name", "") or "").lower()
            if level and text:
                lines.append(f"{'#' * level} {text}")
                heading_count += 1
            elif ("list" in style_name or "列表" in style_name) and text:
                lines.append(f"- {text}")
                list_count += 1
            elif text:
                lines.append(text)
            lines.extend(image_lines)
            lines.append("")
            paragraph_count += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            markdown_table = _docx_table_markdown(table)
            if markdown_table:
                lines.extend([markdown_table, ""])
                table_count += 1

    markdown = "\n".join(lines).strip() + "\n"
    warnings = []
    if not heading_count:
        warnings.append("未识别到 Word 标题样式，导入后需人工校对目录层级")
    if not paragraph_count and not table_count:
        raise DocumentWorkspaceError("Word 文件没有可导入的正文内容")
    return {
        "markdown": markdown,
        "assets": assets,
        "stats": {
            "paragraph_count": paragraph_count,
            "heading_count": heading_count,
            "table_count": table_count,
            "list_count": list_count,
            "image_count": image_count,
            "size_chars": len(markdown),
        },
        "warnings": warnings,
    }


class MultiDocumentService:
    """Keep a project document index while preserving the legacy primary Markdown workspace."""

    def __init__(self) -> None:
        self._presentation_slide_jobs: dict[str, dict[str, Any]] = {}

    def _project_root(self, project: dict[str, Any]) -> Path:
        return document_workspace_service._workspace_root(project)  # noqa: SLF001

    def _index_path(self, project: dict[str, Any]) -> Path:
        return self._project_root(project) / "documents.json"

    def _lock_path(self, project: dict[str, Any]) -> Path:
        return self._project_root(project) / ".documents.lock"

    def _document_root(self, project: dict[str, Any], document_id: str) -> Path:
        root = (self._project_root(project) / "documents" / document_id).resolve()
        try:
            root.relative_to(self._project_root(project).resolve())
        except ValueError as exc:
            raise DocumentWorkspaceError("非法文档路径") from exc
        return root

    def _presentation_manifest_path(self, project: dict[str, Any], document_id: str) -> Path:
        return self._document_root(project, document_id) / "source" / "presentation-manifest.json"

    def _write_json_atomic(self, path: Path, value: dict[str, Any]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(f"{path.suffix}.tmp")
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    def _read_index(self, project: dict[str, Any]) -> dict[str, Any]:
        path = self._index_path(project)
        if not path.exists():
            return {}
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise DocumentWorkspaceError(f"项目文档索引损坏：{path}") from exc
        if not isinstance(data.get("documents"), list):
            raise DocumentWorkspaceError("项目文档索引缺少 documents 数组")
        return data

    def _write_index(self, project: dict[str, Any], index: dict[str, Any]) -> None:
        path = self._index_path(project)
        path.parent.mkdir(parents=True, exist_ok=True)
        index["schema"] = "openclaw.writing-project-documents"
        index["project_id"] = project.get("id")
        index["project_name"] = project.get("name")
        index["updated_at"] = _now()
        temporary = path.with_suffix(".json.tmp")
        temporary.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    def _normalise_record(self, record: dict[str, Any]) -> dict[str, Any]:
        kind = str(record.get("kind") or "")
        is_output_product = bool(record.get("is_output_product", kind == "rich_text"))
        record["is_output_product"] = is_output_product
        default_output_format = (
            "docx"
            if kind == "rich_text" and is_output_product
            else "pptx"
            if kind == "presentation" and is_output_product
            else ""
        )
        record["output_format"] = str(
            record.get("output_format") or default_output_format
        )
        record["data_source_ids"] = [
            str(value)
            for value in (record.get("data_source_ids") or [])
            if str(value).strip()
        ]
        record["product_type"] = str(record.get("product_type") or "")
        record["course_unit_ids"] = list(
            dict.fromkeys(
                str(value)
                for value in (record.get("course_unit_ids") or [])
                if str(value).strip()
            )
        )
        record["quality_profile"] = str(record.get("quality_profile") or "")
        record["required_for_release"] = bool(record.get("required_for_release", False))
        source_refs = []
        for item in record.get("source_refs") or []:
            if not isinstance(item, dict):
                continue
            ref_type = str(item.get("ref_type") or "").strip()
            project_id = str(item.get("project_id") or "").strip()
            document_id = str(item.get("document_id") or "").strip()
            product_id = str(item.get("product_id") or "").strip()
            if not ref_type or not project_id or not (document_id or product_id):
                continue
            source_refs.append(
                {
                    "ref_type": ref_type,
                    "project_id": project_id,
                    "document_id": document_id,
                    "product_id": product_id,
                    "relation": str(item.get("relation") or "").strip(),
                    "required": bool(item.get("required", True)),
                    "version": str(item.get("version") or "").strip(),
                }
            )
        record["source_refs"] = source_refs
        record["print_profile"] = str(
            record.get("print_profile") or ("standard_a4" if is_output_product else "")
        )
        record["publication_status"] = str(
            record.get("publication_status")
            or ("draft" if is_output_product else "internal")
        )
        record["rules_version"] = str(record.get("rules_version") or "")
        record["data_version"] = str(record.get("data_version") or "")
        edit_policy = str(record.get("edit_policy") or "editable").lower()
        record["edit_policy"] = edit_policy if edit_policy in EDIT_POLICIES else "editable"
        delivery_role = str(record.get("delivery_role") or "deliverable").lower()
        record["delivery_role"] = (
            delivery_role if delivery_role in DELIVERY_ROLES else "deliverable"
        )
        lineage = record.get("lineage") if isinstance(record.get("lineage"), dict) else {}
        if lineage:
            source_type = str(lineage.get("source_type") or "markdown").lower()
            try:
                sequence = max(int(lineage.get("sequence") or 0), 0)
            except (TypeError, ValueError):
                sequence = 0
            record["lineage"] = {
                "series_id": str(lineage.get("series_id") or "").strip(),
                "edition_label": str(lineage.get("edition_label") or "").strip(),
                "sequence": sequence,
                "source_type": source_type if source_type in LINEAGE_SOURCE_TYPES else "markdown",
                "parent_document_id": str(lineage.get("parent_document_id") or "").strip(),
                "source_checksum": str(lineage.get("source_checksum") or "").strip().lower(),
                "generated_at": str(lineage.get("generated_at") or "").strip(),
                "source_paths": [
                    str(value) for value in (lineage.get("source_paths") or []) if str(value).strip()
                ],
            }
        else:
            record["lineage"] = None
        try:
            record["expected_chapters"] = max(int(record.get("expected_chapters") or 0), 0)
        except (TypeError, ValueError):
            record["expected_chapters"] = 0
        metadata = record.get("metadata") if isinstance(record.get("metadata"), dict) else {}
        record["metadata"] = metadata
        if isinstance(metadata.get("structure_binding"), dict):
            metadata["structure_binding"] = self._normalise_structure_binding(
                metadata["structure_binding"]
            )
        return record

    def _normalise_structure_binding(self, binding: dict[str, Any] | None) -> dict[str, Any]:
        value = copy.deepcopy(binding) if isinstance(binding, dict) else {}
        mode = str(value.get("mode") or "mapped").lower()
        status = str(value.get("status") or "missing").lower()
        try:
            mapped_items = max(int(value.get("mapped_items") or 0), 0)
        except (TypeError, ValueError):
            mapped_items = 0
        return {
            "mode": mode if mode in STRUCTURE_BINDING_MODES else "mapped",
            "source_document_id": str(value.get("source_document_id") or ""),
            "source_version": str(value.get("source_version") or ""),
            "source_sha256": str(value.get("source_sha256") or "").lower(),
            "status": status if status in STRUCTURE_BINDING_STATUSES else "missing",
            "mapped_items": mapped_items,
            "unmapped_items": [
                _json_value(item) for item in (value.get("unmapped_items") or [])
            ],
            "changed_sections": [
                _json_value(item) for item in (value.get("changed_sections") or [])
            ],
        }

    def _record_checksum(self, project: dict[str, Any], record: dict[str, Any]) -> str:
        kind = str(record.get("kind") or "")
        if kind == "rich_text":
            manifest = document_workspace_service.ensure_workspace(
                self._rich_project(project, record)
            )
            working = Path(str(manifest.get("working_markdown") or ""))
            return hashlib.sha256(working.read_bytes()).hexdigest() if working.is_file() else ""
        path = self._content_path(project, record)
        return hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else ""

    def _record_structure_version(self, record: dict[str, Any]) -> str:
        return str(record.get("data_version") or f"v{int(record.get('revision') or 1)}")

    def _read_presentation_manifest(
        self, project: dict[str, Any], document_id: str
    ) -> dict[str, Any]:
        path = self._presentation_manifest_path(project, document_id)
        if not path.exists():
            return {}
        try:
            value = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise DocumentWorkspaceError("PPT 联动清单损坏") from exc
        if not isinstance(value, dict):
            raise DocumentWorkspaceError("PPT 联动清单格式无效")
        return value

    def _structure_binding(
        self, project: dict[str, Any], record: dict[str, Any]
    ) -> dict[str, Any]:
        metadata = record.get("metadata") if isinstance(record.get("metadata"), dict) else {}
        binding = self._normalise_structure_binding(metadata.get("structure_binding"))
        if not binding["source_document_id"]:
            binding["status"] = "missing"
            return binding

        index = self.ensure_index(project)
        source = next(
            (
                row
                for row in index["documents"]
                if str(row.get("id") or "") == binding["source_document_id"]
            ),
            None,
        )
        if source is None:
            binding["status"] = "missing"
            return binding

        source_sha256 = self._record_checksum(project, source)
        source_version = self._record_structure_version(source)
        stale = (
            not binding["source_sha256"]
            or binding["source_sha256"] != source_sha256
            or (
                bool(binding["source_version"])
                and binding["source_version"] != source_version
            )
        )
        if record.get("kind") == "presentation":
            manifest = self._read_presentation_manifest(project, str(record.get("id") or ""))
            manifest_output = (
                manifest.get("authority", {}).get("presentation_output", {})
                if isinstance(manifest.get("authority"), dict)
                else {}
            )
            expected_ppt_sha256 = str(manifest_output.get("sha256") or "").lower()
            current_ppt_sha256 = self._record_checksum(project, record)
            stale = stale or not expected_ppt_sha256 or expected_ppt_sha256 != current_ppt_sha256

        if stale:
            binding["status"] = "stale"
        elif binding["unmapped_items"] or binding["changed_sections"]:
            binding["status"] = "diverged"
        else:
            binding["status"] = "aligned"
        return binding

    def _validate_presentation_manifest(self, manifest: dict[str, Any]) -> dict[str, Any]:
        value = copy.deepcopy(manifest)
        slides = value.get("slides")
        if not isinstance(slides, list):
            raise DocumentWorkspaceError("PPT 联动清单缺少 slides 数组")
        slide_numbers = []
        for item in slides:
            if not isinstance(item, dict):
                raise DocumentWorkspaceError("PPT 联动清单包含无效页面")
            try:
                slide_number = int(item.get("slide"))
            except (TypeError, ValueError) as exc:
                raise DocumentWorkspaceError("PPT 联动清单页码无效") from exc
            if slide_number < 1:
                raise DocumentWorkspaceError("PPT 联动清单页码无效")
            slide_numbers.append(slide_number)
        if len(slide_numbers) != len(set(slide_numbers)):
            raise DocumentWorkspaceError("PPT 联动清单存在重复页码映射")
        expected = list(range(1, len(slide_numbers) + 1))
        if sorted(slide_numbers) != expected:
            raise DocumentWorkspaceError("PPT 联动清单页码必须从 1 连续编号")
        return value

    def ensure_index(self, project: dict[str, Any], primary_title: str | None = None) -> dict[str, Any]:
        self._project_root(project).mkdir(parents=True, exist_ok=True)
        with self._lock_path(project).open("w", encoding="utf-8") as lock:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
            index = self._read_index(project)
            if index:
                previous = json.dumps(index.get("documents") or [], ensure_ascii=False, sort_keys=True)
                for record in index["documents"]:
                    self._normalise_record(record)
                current = json.dumps(index.get("documents") or [], ensure_ascii=False, sort_keys=True)
                if index.get("project_name") != project.get("name") or previous != current:
                    self._write_index(project, index)
                return copy.deepcopy(index)
            manifest = document_workspace_service.ensure_workspace(project)
            primary_key = f"{project.get('id')}:primary"
            document_id = f"doc-{hashlib.sha1(primary_key.encode()).hexdigest()[:12]}"
            title = primary_title or self._default_primary_title(project)
            created_at = manifest.get("created_at") or _now()
            index = {
                "schema": "openclaw.writing-project-documents",
                "project_id": project.get("id"),
                "project_name": project.get("name"),
                "created_at": created_at,
                "documents": [{
                    "id": document_id,
                    "title": title,
                    "kind": "rich_text",
                    "status": "active",
                    "sort_order": 0,
                    "is_primary": True,
                    "legacy_primary": True,
                    "revision": int(manifest.get("version") or 1),
                    "source_path": manifest.get("source_markdown") or "",
                    "created_at": created_at,
                    "updated_at": manifest.get("updated_at") or created_at,
                    "metadata": {},
                    "is_output_product": True,
                    "output_format": "docx",
                    "data_source_ids": [],
                    "print_profile": "standard_a4",
                    "publication_status": "draft",
                    "rules_version": "",
                    "data_version": "",
                    "expected_chapters": max(
                        int(manifest.get("stats", {}).get("chapter_count") or 0),
                        int(project.get("document_spec", {}).get("expected_chapters") or 0),
                    ),
                }],
            }
            self._write_index(project, index)
            return copy.deepcopy(index)

    def _default_primary_title(self, project: dict[str, Any]) -> str:
        name = str(project.get("name") or "")
        if "博士论文" in name:
            return "博士论文正文"
        if "兵棋" in name or "规则手册" in name:
            return "规则手册"
        return "正文"

    def _get_record(self, project: dict[str, Any], document_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
        index = self.ensure_index(project)
        record = next((row for row in index["documents"] if row.get("id") == document_id), None)
        if not record:
            raise DocumentWorkspaceError("文档不存在")
        return index, record

    def _rich_project(self, project: dict[str, Any], record: dict[str, Any]) -> dict[str, Any]:
        if record.get("legacy_primary"):
            synthetic = copy.deepcopy(project)
        else:
            root = self._document_root(project, record["id"])
            source = root / "source" / "document.md"
            synthetic = copy.deepcopy(project)
            synthetic["name"] = record.get("title") or project.get("name")
            synthetic["_workspace_root_override"] = str(root)
            spec = synthetic.get("document_spec") if isinstance(synthetic.get("document_spec"), dict) else {}
            spec = copy.deepcopy(spec)
            working_spec = spec.get("working_markdown") if isinstance(spec.get("working_markdown"), dict) else {}
            working_spec = {**working_spec, "path": str(source)}
            spec["working_markdown"] = working_spec
            source_word = root / "source" / "original.docx"
            if source_word.is_file():
                spec["source_word"] = {"path": str(source_word)}
            synthetic["document_spec"] = spec
        expected_chapters = int(record.get("expected_chapters") or 0)
        if expected_chapters:
            spec = synthetic.get("document_spec") if isinstance(synthetic.get("document_spec"), dict) else {}
            spec = copy.deepcopy(spec)
            spec["expected_chapters"] = expected_chapters
            synthetic["document_spec"] = spec
        synthetic["_document_title"] = record.get("title") or project.get("name") or "正文文档"
        synthetic["_document_id"] = record.get("id") or ""
        synthetic["_document_print_profile"] = record.get("print_profile") or "standard_a4"
        synthetic["_course_document_record"] = self._normalise_record(copy.deepcopy(record))
        metadata = record.get("metadata") if isinstance(record.get("metadata"), dict) else {}
        synthetic["_document_layout_binding"] = copy.deepcopy(metadata.get("layout_binding") or {})
        return synthetic

    def list_documents(self, project: dict[str, Any], include_archived: bool = True) -> dict[str, Any]:
        index = self.ensure_index(project)
        rows = []
        for record in sorted(index["documents"], key=lambda row: (int(row.get("sort_order") or 0), row.get("title") or "")):
            if not include_archived and record.get("status") == "archived":
                continue
            rows.append(self._public_record(project, record))
        return {
            "project": {
                "id": project.get("id"),
                "name": project.get("name"),
                "description": project.get("description"),
                "status": project.get("status"),
                "progress": project.get("progress"),
            },
            "documents": rows,
            "summary": {
                "total": len(rows),
                "active": sum(1 for row in rows if row["status"] == "active"),
                "archived": sum(1 for row in rows if row["status"] == "archived"),
                "rich_text": sum(1 for row in rows if row["kind"] == "rich_text"),
                "workbook": sum(1 for row in rows if row["kind"] == "workbook"),
                "presentation": sum(1 for row in rows if row["kind"] == "presentation"),
                "output_products": sum(1 for row in rows if row["is_output_product"]),
                "internal_sources": sum(1 for row in rows if not row["is_output_product"]),
            },
        }

    def _public_record(self, project: dict[str, Any], record: dict[str, Any]) -> dict[str, Any]:
        row = self._normalise_record(copy.deepcopy(record))
        kind = row.get("kind")
        try:
            if kind == "rich_text":
                manifest = document_workspace_service.ensure_workspace(self._rich_project(project, row))
                row["revision"] = int(manifest.get("version") or row.get("revision") or 1)
                row["stats"] = manifest.get("stats") or {}
                row["updated_at"] = manifest.get("updated_at") or row.get("updated_at")
                working = Path(str(manifest.get("working_markdown") or ""))
                if working.is_file():
                    row["source_checksum"] = hashlib.sha256(working.read_bytes()).hexdigest()
            elif kind == "workbook" and self._content_path(project, row).exists():
                summary = self._workbook_summary(self._content_path(project, row))
                row["stats"] = summary
            elif kind == "presentation" and self._content_path(project, row).exists():
                presentation_manifest = self._read_presentation_manifest(
                    project, str(row.get("id") or "")
                )
                presentation_output = (
                    presentation_manifest.get("authority", {}).get(
                        "presentation_output", {}
                    )
                    if isinstance(presentation_manifest.get("authority"), dict)
                    else {}
                )
                row["stats"] = {
                    "slide_count": self._presentation_slide_count(self._content_path(project, row)),
                    "preview_ready": self._presentation_pdf_path(project, row).exists(),
                    "main_slide_count": int(
                        presentation_output.get("main_slide_count") or 0
                    ),
                    "appendix_slide_count": int(
                        presentation_output.get("appendix_slide_count") or 0
                    ),
                    "notes_count": int(presentation_output.get("notes_count") or 0),
                }
            else:
                row["stats"] = {}
        except (DocumentWorkspaceError, OSError, BadZipFile, ValueError):
            row["health"] = "error"
            row["stats"] = {}
        row["structure_binding"] = self._structure_binding(project, row)
        row["source_path"] = str(row.get("source_path") or "")
        return row

    def create_document(
        self,
        project: dict[str, Any],
        title: str,
        kind: str,
        *,
        outline: list[str] | None = None,
        source_path: str = "",
        is_primary: bool = False,
        is_output_product: bool | None = None,
        output_format: str = "",
        data_source_ids: list[str] | None = None,
        product_type: str = "",
        course_unit_ids: list[str] | None = None,
        quality_profile: str = "",
        required_for_release: bool = False,
        source_refs: list[dict[str, Any]] | None = None,
        print_profile: str = "",
        publication_status: str = "",
        rules_version: str = "",
        data_version: str = "",
        edit_policy: str = "editable",
        delivery_role: str = "deliverable",
        lineage: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        title = title.strip()
        if not title:
            raise DocumentWorkspaceError("文档名称不能为空")
        if kind not in DOCUMENT_KINDS:
            raise DocumentWorkspaceError("不支持的文档类型")
        output_product = kind == "rich_text" if is_output_product is None else bool(is_output_product)
        if output_product and kind not in {"rich_text", "presentation"}:
            raise DocumentWorkspaceError("正式输出产品仅支持正文或演示文档")
        selected_format = output_format.strip().lower() or (
            "docx"
            if output_product and kind == "rich_text"
            else "pptx"
            if output_product and kind == "presentation"
            else ""
        )
        if selected_format and selected_format not in OUTPUT_FORMATS:
            raise DocumentWorkspaceError("正式输出格式只支持 docx、pdf 或 pptx")
        if output_product and (
            (kind == "rich_text" and selected_format not in {"docx", "pdf"})
            or (kind == "presentation" and selected_format != "pptx")
        ):
            raise DocumentWorkspaceError("正式输出格式与文档类型不匹配")
        selected_status = publication_status.strip().lower() or (
            "draft" if output_product else "internal"
        )
        if selected_status not in PUBLICATION_STATUSES:
            raise DocumentWorkspaceError("无效发布状态")
        selected_edit_policy = edit_policy.strip().lower() or "editable"
        if selected_edit_policy not in EDIT_POLICIES:
            raise DocumentWorkspaceError("无效编辑策略")
        selected_delivery_role = delivery_role.strip().lower() or "deliverable"
        if selected_delivery_role not in DELIVERY_ROLES:
            raise DocumentWorkspaceError("无效交付角色")
        self.ensure_index(project)
        with self._lock_path(project).open("w", encoding="utf-8") as lock:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
            index = self._read_index(project)
            available_ids = {str(row.get("id") or "") for row in index["documents"]}
            requested_source_ids = [
                str(value) for value in (data_source_ids or []) if str(value).strip()
            ]
            if any(value not in available_ids for value in requested_source_ids):
                raise DocumentWorkspaceError("关联数据源不存在")
            if any(str(row.get("title") or "").casefold() == title.casefold() and row.get("status") != "archived" for row in index["documents"]):
                raise DocumentWorkspaceError("项目内已存在同名文档")
            document_id = f"doc-{uuid.uuid4().hex[:12]}"
            now = _now()
            chapter_outline = [item.strip() for item in (outline or []) if item.strip()]
            record = {
                "id": document_id,
                "title": title,
                "kind": kind,
                "status": "active",
                "sort_order": max([int(row.get("sort_order") or 0) for row in index["documents"]] + [-1]) + 1,
                "is_primary": bool(is_primary),
                "legacy_primary": False,
                "revision": 1 if kind == "rich_text" or source_path else 0,
                "source_path": "",
                "created_at": now,
                "updated_at": now,
                "metadata": {},
                "is_output_product": output_product,
                "output_format": selected_format,
                "data_source_ids": list(dict.fromkeys(requested_source_ids)),
                "product_type": product_type.strip(),
                "course_unit_ids": list(
                    dict.fromkeys(
                        str(value)
                        for value in (course_unit_ids or [])
                        if str(value).strip()
                    )
                ),
                "quality_profile": quality_profile.strip(),
                "required_for_release": bool(required_for_release),
                "source_refs": copy.deepcopy(source_refs or []),
                "print_profile": print_profile.strip() or (
                    "standard_a4" if output_product else ""
                ),
                "publication_status": selected_status,
                "rules_version": rules_version.strip(),
                "data_version": data_version.strip(),
                "expected_chapters": (
                    len(chapter_outline)
                    if kind == "rich_text" and chapter_outline
                    else 3 if kind == "rich_text" else 0
                ),
                "edit_policy": selected_edit_policy,
                "delivery_role": selected_delivery_role,
                "lineage": copy.deepcopy(lineage) if lineage else None,
            }
            self._normalise_record(record)
            if is_primary:
                for row in index["documents"]:
                    row["is_primary"] = False
            index["documents"].append(record)
            self._write_index(project, index)
        root = self._document_root(project, document_id)
        for name in ["source", "working", "versions", "exports", "rendered"]:
            (root / name).mkdir(parents=True, exist_ok=True)
        if source_path:
            source = self._validated_vault_source(source_path)
            self._install_source(project, record, source)
        elif kind == "rich_text":
            self._initialize_rich_text(project, record, outline or [])
        return self.get_document(project, document_id)

    def _validated_vault_source(self, source_path: str) -> Path:
        source = Path(source_path).expanduser().resolve()
        try:
            source.relative_to(document_workspace_service.vault)
        except ValueError as exc:
            raise DocumentWorkspaceError("导入源文件必须位于知识库内") from exc
        if not source.is_file():
            raise DocumentWorkspaceError("导入源文件不存在")
        return source

    def _initialize_rich_text(self, project: dict[str, Any], record: dict[str, Any], outline: list[str]) -> None:
        root = self._document_root(project, record["id"])
        source = root / "source" / "document.md"
        chapters = [item.strip() for item in outline if item.strip()] or ["第一章 背景与目标", "第二章 核心内容", "第三章 总结"]
        lines = [
            f"# {record['title']}",
            "",
            f"**所属项目：** {project.get('name') or ''}",
            "",
        ]
        for chapter in chapters:
            lines.extend([f"# {chapter}", "", "## 本章要点", "", "待撰写。", ""])
        lines.extend(["# 参考文献", "", "待补充。", ""])
        source.write_text("\n".join(lines), encoding="utf-8")
        manifest = document_workspace_service.ensure_workspace(self._rich_project(project, record))
        self._update_record(project, record["id"], {
            "source_path": str(source),
            "revision": int(manifest.get("version") or 1),
            "updated_at": manifest.get("updated_at") or _now(),
        })

    def _install_source(self, project: dict[str, Any], record: dict[str, Any], source: Path) -> None:
        expected = {"rich_text": ".md", "workbook": ".xlsx", "presentation": ".pptx"}[record["kind"]]
        if source.suffix.lower() != expected:
            raise DocumentWorkspaceError(f"{record['kind']} 文档只支持 {expected} 文件")
        root = self._document_root(project, record["id"])
        source_copy = root / "source" / f"original{expected}"
        source_copy.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, source_copy)
        if record["kind"] == "rich_text":
            shutil.copy2(source_copy, root / "source" / "document.md")
            document_workspace_service.ensure_workspace(self._rich_project(project, record))
        else:
            working = self._content_path(project, record)
            working.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source_copy, working)
            shutil.copy2(working, root / "versions" / f"v0001-import{expected}")
            if record["kind"] == "presentation":
                self._render_presentation(project, record)
        self._update_record(project, record["id"], {
            "source_path": str(source),
            "source_checksum": hashlib.sha256(source.read_bytes()).hexdigest(),
            "updated_at": _now(),
        })

    def import_docx(
        self,
        project: dict[str, Any],
        title: str,
        data: bytes,
        filename: str,
        *,
        is_primary: bool = False,
        product_type: str = "",
        course_unit_ids: list[str] | None = None,
        quality_profile: str = "",
        required_for_release: bool = False,
        source_refs: list[dict[str, Any]] | None = None,
        actor: str = "human-editor",
    ) -> dict[str, Any]:
        if len(data) > MAX_UPLOAD_BYTES:
            raise DocumentWorkspaceError("文件超过 100 MB 上限")
        if Path(filename).suffix.lower() != ".docx":
            raise DocumentWorkspaceError("只支持 .docx 文件")
        with ZipFile(io.BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise DocumentWorkspaceError("Word 文件结构无效")
        converted = _convert_docx(data, title)
        document = self.create_document(
            project,
            title,
            "rich_text",
            is_primary=is_primary,
            product_type=product_type,
            course_unit_ids=course_unit_ids,
            quality_profile=quality_profile,
            required_for_release=required_for_release,
            source_refs=source_refs,
        )
        document_id = str(document["id"])
        _, record = self._get_record(project, document_id)
        root = self._document_root(project, document_id)
        source = root / "source" / "document.md"
        original = root / "source" / "original.docx"
        working = root / "working" / "document.md"
        versions = root / "versions"
        assets_dir = root / "source" / "assets"
        versions.mkdir(parents=True, exist_ok=True)
        assets_dir.mkdir(parents=True, exist_ok=True)
        if working.exists():
            shutil.copy2(working, versions / "v0001-before-docx-import.md")
        original.write_bytes(data)
        source.write_text(converted["markdown"], encoding="utf-8")
        working.write_text(converted["markdown"], encoding="utf-8")
        for asset_name, asset_data in converted["assets"].items():
            (assets_dir / asset_name).write_bytes(asset_data)

        synthetic = self._rich_project(project, record)
        manifest = document_workspace_service.ensure_workspace(synthetic)
        markdown = converted["markdown"]
        sections = document_workspace_service._parse_sections(markdown)  # noqa: SLF001
        manifest.update(
            {
                "version": max(int(manifest.get("version") or 1) + 1, 2),
                "source_markdown": str(source),
                "source_base_dir": str(source.parent),
                "source_word": str(original),
                "working_markdown": str(working),
                "stats": document_workspace_service._stats(markdown, sections),  # noqa: SLF001
            }
        )
        document_workspace_service._write_manifest(synthetic, manifest)  # noqa: SLF001
        self._update_record(
            project,
            document_id,
            {
                "source_path": str(source),
                "source_checksum": hashlib.sha256(data).hexdigest(),
                "revision": int(manifest["version"]),
                "updated_at": _now(),
                "metadata": {
                    **(record.get("metadata") or {}),
                    "original_filename": Path(filename).name,
                    "import_actor": actor,
                    "docx_import": {
                        **converted["stats"],
                        "warnings": converted["warnings"],
                    },
                },
            },
        )
        result = self.get_document(project, document_id)
        result["import_summary"] = {
            **converted["stats"],
            "warnings": converted["warnings"],
            "original_retained": True,
        }
        return result

    def get_document(self, project: dict[str, Any], document_id: str) -> dict[str, Any]:
        _, record = self._get_record(project, document_id)
        return self._public_record(project, record)

    def rich_project_context(self, project: dict[str, Any], document_id: str) -> dict[str, Any]:
        """Return the generic rich-text project context used by document services."""
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "rich_text":
            raise DocumentWorkspaceError("该操作仅支持正文文档")
        return self._rich_project(project, record)

    def assert_writable(self, project: dict[str, Any], document_id: str) -> dict[str, Any]:
        """Reject every content mutation for immutable historical documents."""
        _, record = self._get_record(project, document_id)
        self._normalise_record(record)
        if record.get("edit_policy") != "editable":
            raise DocumentVersionConflict("历史版本为只读文档，不能修改正文或运行 AI 写入任务")
        return record

    def update_document(self, project: dict[str, Any], document_id: str, patch: dict[str, Any]) -> dict[str, Any]:
        allowed = {
            "title",
            "status",
            "sort_order",
            "is_primary",
            "is_output_product",
            "output_format",
            "data_source_ids",
            "product_type",
            "course_unit_ids",
            "quality_profile",
            "required_for_release",
            "source_refs",
            "print_profile",
            "publication_status",
            "rules_version",
            "data_version",
            "expected_chapters",
            "edit_policy",
            "delivery_role",
            "lineage",
        }
        changes = {key: value for key, value in patch.items() if key in allowed}
        if "title" in changes:
            changes["title"] = str(changes["title"]).strip()
            if not changes["title"]:
                raise DocumentWorkspaceError("文档名称不能为空")
        if "status" in changes and changes["status"] not in {"active", "archived"}:
            raise DocumentWorkspaceError("无效文档状态")
        index, record = self._get_record(project, document_id)
        self._normalise_record(record)
        output_product = bool(changes.get("is_output_product", record.get("is_output_product")))
        if output_product and record.get("kind") not in {"rich_text", "presentation"}:
            raise DocumentWorkspaceError("正式输出产品仅支持正文或演示文档")
        if "output_format" in changes:
            changes["output_format"] = str(changes["output_format"]).strip().lower()
            if changes["output_format"] and changes["output_format"] not in OUTPUT_FORMATS:
                raise DocumentWorkspaceError("正式输出格式只支持 docx、pdf 或 pptx")
        selected_format = str(changes.get("output_format", record.get("output_format") or ""))
        if output_product and (
            (record.get("kind") == "rich_text" and selected_format not in {"docx", "pdf"})
            or (record.get("kind") == "presentation" and selected_format != "pptx")
        ):
            raise DocumentWorkspaceError("正式输出格式与文档类型不匹配")
        if "publication_status" in changes:
            changes["publication_status"] = str(changes["publication_status"]).strip().lower()
            if changes["publication_status"] not in PUBLICATION_STATUSES:
                raise DocumentWorkspaceError("无效发布状态")
        if "edit_policy" in changes:
            changes["edit_policy"] = str(changes["edit_policy"]).strip().lower()
            if changes["edit_policy"] not in EDIT_POLICIES:
                raise DocumentWorkspaceError("无效编辑策略")
        if "delivery_role" in changes:
            changes["delivery_role"] = str(changes["delivery_role"]).strip().lower()
            if changes["delivery_role"] not in DELIVERY_ROLES:
                raise DocumentWorkspaceError("无效交付角色")
        if "lineage" in changes:
            temporary = {"lineage": changes["lineage"]}
            self._normalise_record(temporary)
            changes["lineage"] = temporary["lineage"]
        if "data_source_ids" in changes:
            source_ids = [
                str(value) for value in (changes["data_source_ids"] or []) if str(value).strip()
            ]
            available_ids = {str(row.get("id") or "") for row in index["documents"]}
            if document_id in source_ids or any(value not in available_ids for value in source_ids):
                raise DocumentWorkspaceError("关联数据源不存在或形成自引用")
            changes["data_source_ids"] = list(dict.fromkeys(source_ids))
        if "course_unit_ids" in changes:
            changes["course_unit_ids"] = list(
                dict.fromkeys(
                    str(value)
                    for value in (changes["course_unit_ids"] or [])
                    if str(value).strip()
                )
            )
        if "source_refs" in changes:
            temporary = {"source_refs": changes["source_refs"]}
            self._normalise_record(temporary)
            changes["source_refs"] = temporary["source_refs"]
        if "required_for_release" in changes:
            changes["required_for_release"] = bool(changes["required_for_release"])
        for key in {"print_profile", "rules_version", "data_version", "product_type", "quality_profile"}:
            if key in changes:
                changes[key] = str(changes[key]).strip()
        if "expected_chapters" in changes:
            try:
                changes["expected_chapters"] = max(int(changes["expected_chapters"]), 0)
            except (TypeError, ValueError) as exc:
                raise DocumentWorkspaceError("计划章节数必须是非负整数") from exc
        if changes.get("status") == "archived" and record.get("is_primary"):
            raise DocumentWorkspaceError("主文档不能归档，请先设置其他主文档")
        if changes.get("is_primary"):
            for row in index["documents"]:
                row["is_primary"] = row.get("id") == document_id
        record.update(changes)
        record["updated_at"] = _now()
        self._write_index(project, index)
        return self._public_record(project, record)

    def _update_record(self, project: dict[str, Any], document_id: str, changes: dict[str, Any]) -> None:
        index, record = self._get_record(project, document_id)
        record.update(changes)
        self._write_index(project, index)

    def update_document_metadata(
        self,
        project: dict[str, Any],
        document_id: str,
        patch: dict[str, Any],
    ) -> dict[str, Any]:
        """Merge system-owned metadata without exposing arbitrary record fields."""
        index, record = self._get_record(project, document_id)
        metadata = record.get("metadata") if isinstance(record.get("metadata"), dict) else {}
        record["metadata"] = {**metadata, **copy.deepcopy(patch)}
        record["updated_at"] = _now()
        self._write_index(project, index)
        return self._public_record(project, record)

    def detach_legacy_primary(self, project: dict[str, Any], document_id: str) -> dict[str, Any]:
        index, record = self._get_record(project, document_id)
        if not record.get("legacy_primary"):
            return self._public_record(project, record)
        legacy_manifest = document_workspace_service.ensure_workspace(project)
        legacy_working = Path(str(legacy_manifest.get("working_markdown") or ""))
        if not legacy_working.is_file():
            raise DocumentWorkspaceError("兼容主文档工作稿不存在")
        root = self._document_root(project, document_id)
        for name in ["source", "working", "versions", "exports", "rendered"]:
            (root / name).mkdir(parents=True, exist_ok=True)
        source = root / "source" / "document.md"
        shutil.copy2(legacy_working, source)
        shutil.copy2(legacy_working, root / "versions" / "v0001-legacy-primary.md")
        source_word = Path(str(legacy_manifest.get("source_word") or ""))
        if source_word.is_file():
            shutil.copy2(source_word, root / "source" / "original.docx")
        record["legacy_primary"] = False
        record["source_path"] = str(source)
        record["updated_at"] = _now()
        self._write_index(project, index)
        document_workspace_service.ensure_workspace(self._rich_project(project, record))
        return self._public_record(project, record)

    def replace_rich_text_markdown(
        self,
        project: dict[str, Any],
        document_id: str,
        markdown: str,
        actor: str,
        target_version: int | None = None,
    ) -> dict[str, Any]:
        self.assert_writable(project, document_id)
        index, record = self._get_record(project, document_id)
        if record.get("kind") != "rich_text":
            raise DocumentWorkspaceError("当前文档不是正文文档")
        if not markdown.strip():
            raise DocumentWorkspaceError("正文内容不能为空")
        synthetic = self._rich_project(project, record)
        manifest = document_workspace_service.ensure_workspace(synthetic)
        if manifest.get("content_authority") == "structured_json":
            raise DocumentWorkspaceError(
                "该文档已启用结构化协同编辑，旧 Markdown 替换接口仅提供只读兼容"
            )
        working = Path(str(manifest.get("working_markdown") or ""))
        source = Path(str(manifest.get("source_markdown") or ""))
        versions = document_workspace_service._workspace_root(synthetic) / "versions"  # noqa: SLF001
        versions.mkdir(parents=True, exist_ok=True)
        current_version = max(int(manifest.get("version") or 1), 1)
        if working.is_file():
            shutil.copy2(
                working,
                versions / f"v{current_version:04d}-before-{_safe_name(actor)}.md",
            )
        next_version = current_version + 1
        if target_version is not None:
            try:
                requested_version = int(target_version)
            except (TypeError, ValueError) as exc:
                raise DocumentWorkspaceError("目标正文版本必须是整数") from exc
            if requested_version <= current_version:
                raise DocumentWorkspaceError(
                    f"目标正文版本必须高于当前版本 v{current_version}"
                )
            next_version = requested_version
        working.write_text(markdown, encoding="utf-8")
        if source.is_file() and source.resolve() != working.resolve():
            source.write_text(markdown, encoding="utf-8")
        sections = document_workspace_service._parse_sections(markdown)  # noqa: SLF001
        manifest["version"] = next_version
        manifest["stats"] = document_workspace_service._stats(markdown, sections)  # noqa: SLF001
        document_workspace_service._write_manifest(synthetic, manifest)  # noqa: SLF001
        record["revision"] = manifest["version"]
        record["source_path"] = str(source)
        record["source_checksum"] = hashlib.sha256(markdown.encode("utf-8")).hexdigest()
        record["updated_at"] = _now()
        record.setdefault("metadata", {})["last_actor"] = actor
        for dependent in index["documents"]:
            binding = (dependent.get("metadata") or {}).get("structure_binding")
            if isinstance(binding, dict) and binding.get("source_document_id") == document_id:
                binding["status"] = "stale"
        self._write_index(project, index)
        return self._public_record(project, record)

    def reorder_documents(self, project: dict[str, Any], document_ids: list[str]) -> list[dict[str, Any]]:
        index = self.ensure_index(project)
        existing = {row["id"]: row for row in index["documents"]}
        if set(document_ids) != set(existing):
            raise DocumentWorkspaceError("排序列表必须包含项目内全部文档")
        for order, document_id in enumerate(document_ids):
            existing[document_id]["sort_order"] = order
        self._write_index(project, index)
        return self.list_documents(project)["documents"]

    def delete_document(self, project: dict[str, Any], document_id: str) -> None:
        index, record = self._get_record(project, document_id)
        if record.get("status") != "archived":
            raise DocumentWorkspaceError("只能永久删除已归档文档")
        if record.get("is_primary") or record.get("legacy_primary"):
            raise DocumentWorkspaceError("主正文及兼容工作区不能永久删除")
        root = self._document_root(project, document_id)
        index["documents"] = [row for row in index["documents"] if row.get("id") != document_id]
        self._write_index(project, index)
        if root.exists():
            shutil.rmtree(root)

    def rich_workspace(self, project: dict[str, Any], document_id: str) -> dict[str, Any]:
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "rich_text":
            raise DocumentWorkspaceError("当前文档不是正文文档")
        workspace = document_workspace_service.workspace(self._rich_project(project, record))
        workspace["document"] = self._public_record(project, record)
        workspace["project"]["id"] = project.get("id")
        workspace["project"]["name"] = project.get("name")
        workspace["project"]["description"] = project.get("description")
        return workspace

    def rich_project(self, project: dict[str, Any], document_id: str) -> dict[str, Any]:
        """Return the isolated project context used by a rich-text document."""
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "rich_text":
            raise DocumentWorkspaceError("当前文档不是正文文档")
        return self._rich_project(project, record)

    def rich_call(self, project: dict[str, Any], document_id: str, method: str, *args: Any) -> Any:
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "rich_text":
            raise DocumentWorkspaceError("当前文档不是正文文档")
        target = getattr(document_workspace_service, method)
        return target(self._rich_project(project, record), *args)

    def upload_rich_text_asset(
        self,
        project: dict[str, Any],
        document_id: str,
        data: bytes,
        filename: str,
        content_type: str,
        actor: str = "human-editor",
    ) -> dict[str, Any]:
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "rich_text":
            raise DocumentWorkspaceError("当前文档不是正文文档")
        if not data:
            raise DocumentWorkspaceError("上传文件不能为空")
        if len(data) > MAX_ASSET_UPLOAD_BYTES:
            raise DocumentWorkspaceError("图片超过 20 MB 上限")
        media_type = content_type.split(";", 1)[0].strip().lower()
        if media_type not in IMAGE_ASSET_TYPES:
            raise DocumentWorkspaceError("只支持 png、jpg、webp、gif 或 svg 图片")
        original_stem = _safe_name(Path(filename or "image").stem or "image")
        suffix = Path(filename or "").suffix.lower() or IMAGE_ASSET_TYPES[media_type]
        if suffix == ".jpeg":
            suffix = ".jpg"
        if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg"}:
            suffix = IMAGE_ASSET_TYPES[media_type]
        root = self._document_root(project, document_id)
        assets_dir = root / "source" / "assets"
        assets_dir.mkdir(parents=True, exist_ok=True)
        digest = hashlib.sha256(data).hexdigest()[:12]
        target_name = f"{original_stem}-{digest}{suffix}"
        target = assets_dir / target_name
        target.write_bytes(data)
        synthetic = self._rich_project(project, record)
        manifest = document_workspace_service.ensure_workspace(synthetic)
        metadata = record.get("metadata") if isinstance(record.get("metadata"), dict) else {}
        uploads = list(metadata.get("asset_uploads") or [])
        uploads.append({
            "path": f"assets/{target_name}",
            "filename": target_name,
            "content_type": media_type,
            "size_bytes": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
            "actor": actor,
            "created_at": _now(),
        })
        metadata["asset_uploads"] = uploads[-100:]
        self._update_record(project, document_id, {
            "metadata": metadata,
            "updated_at": _now(),
        })
        return {
            "path": f"assets/{target_name}",
            "filename": target_name,
            "content_type": media_type,
            "size_bytes": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
            "url": (
                f"/api/v3/writing/projects/{project.get('id')}/documents/{document_id}"
                f"/assets?path=assets/{target_name}"
            ),
            "source_base_dir": manifest.get("source_base_dir"),
        }

    def export_package(self, project: dict[str, Any]) -> Path:
        profile = (
            (project.get("document_spec") or {}).get("course_profile")
            if isinstance(project.get("document_spec"), dict)
            else {}
        )
        if isinstance(profile, dict) and profile.get("template_key"):
            from services.course_production_service import course_production_service

            course_production_service.require_export_ready(project)
        index = self.ensure_index(project)
        records = [
            self._normalise_record(copy.deepcopy(record))
            for record in sorted(
                index["documents"],
                key=lambda row: (int(row.get("sort_order") or 0), row.get("title") or ""),
            )
            if record.get("status") == "active"
            and bool(record.get("is_output_product", record.get("kind") == "rich_text"))
            and str(record.get("delivery_role") or "deliverable") == "deliverable"
        ]
        if not records:
            raise DocumentWorkspaceError("项目没有可交付的正式文档")
        invalid = [
            record.get("title") or record.get("id")
            for record in records
            if (
                record.get("kind") == "rich_text"
                and record.get("output_format") not in {"docx", "pdf"}
            )
            or (
                record.get("kind") == "presentation"
                and record.get("output_format") != "pptx"
            )
            or record.get("kind") not in {"rich_text", "presentation"}
            or record.get("publication_status") not in {"approved", "published"}
        ]
        if invalid:
            raise DocumentWorkspaceError(
                f"正式文档尚未批准或输出格式无效：{'、'.join(str(value) for value in invalid)}"
            )

        exports = self._project_root(project) / "exports"
        exports.mkdir(parents=True, exist_ok=True)
        project_name = _safe_name(str(project.get("name") or "文档项目"))
        package = exports / f"{project_name}-正式交付包.zip"
        temporary = package.with_suffix(".zip.tmp")
        arc_names: set[str] = set()
        with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as archive:
            for record in records:
                if record.get("kind") == "rich_text":
                    output_format = str(record.get("output_format") or "docx")
                    exported = document_workspace_service.export(
                        self._rich_project(project, record),
                        output_format,
                    )
                else:
                    output_format = "pptx"
                    exported = self._content_path(project, record)
                arc_name = (
                    f"{project_name}——{_safe_name(str(record.get('title') or '文档'))}"
                    f".{output_format}"
                )
                if arc_name in arc_names:
                    raise DocumentWorkspaceError("正式文档名称重复，无法生成交付包")
                arc_names.add(arc_name)
                archive.write(exported, arc_name)
        temporary.replace(package)
        return package

    def _content_path(self, project: dict[str, Any], record: dict[str, Any]) -> Path:
        suffix = ".xlsx" if record.get("kind") == "workbook" else ".pptx"
        name = "workbook" if record.get("kind") == "workbook" else "presentation"
        return self._document_root(project, record["id"]) / "working" / f"{name}{suffix}"

    def _workbook_summary(self, path: Path) -> dict[str, Any]:
        workbook = load_workbook(path, read_only=False, data_only=False)
        formulas = 0
        cells = 0
        sheets = []
        for sheet in workbook.worksheets:
            sheet_formulas = 0
            sheet_cells = 0
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        sheet_cells += 1
                        if cell.data_type == "f":
                            sheet_formulas += 1
            formulas += sheet_formulas
            cells += sheet_cells
            sheets.append({
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "formula_count": sheet_formulas,
                "cell_count": sheet_cells,
            })
        return {"sheet_count": len(sheets), "formula_count": formulas, "cell_count": cells, "sheets": sheets}

    def workbook_metadata(self, project: dict[str, Any], document_id: str) -> dict[str, Any]:
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "workbook":
            raise DocumentWorkspaceError("当前文档不是表格文档")
        path = self._content_path(project, record)
        if not path.exists():
            raise DocumentWorkspaceError("尚未上传 XLSX 文件")
        return {
            "document": self._public_record(project, record),
            **self._workbook_summary(path),
        }

    def workbook_sheet(self, project: dict[str, Any], document_id: str, sheet_name: str) -> dict[str, Any]:
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "workbook":
            raise DocumentWorkspaceError("当前文档不是表格文档")
        path = self._content_path(project, record)
        workbook = load_workbook(path, read_only=False, data_only=False)
        values = load_workbook(path, read_only=False, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise DocumentWorkspaceError("工作表不存在")
        sheet = workbook[sheet_name]
        value_sheet = values[sheet_name]
        max_row = min(max(sheet.max_row, 1), 5000)
        max_column = min(max(sheet.max_column, 1), 200)
        cells = []
        for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                display_value = value_sheet[cell.coordinate].value if cell.data_type == "f" else cell.value
                cells.append({
                    "coordinate": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "value": _json_value(cell.value),
                    "display": _json_value(display_value),
                    "formula": str(cell.value) if cell.data_type == "f" else "",
                    "data_type": cell.data_type,
                    "number_format": cell.number_format,
                    "style": {
                        "bold": bool(cell.font.bold),
                        "italic": bool(cell.font.italic),
                        "horizontal": cell.alignment.horizontal or "",
                        "vertical": cell.alignment.vertical or "",
                        "fill": cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor.type == "rgb" else "",
                    },
                })
        validations = []
        if sheet.data_validations:
            for validation in sheet.data_validations.dataValidation:
                validations.append({
                    "ranges": str(validation.sqref),
                    "type": validation.type,
                    "formula1": validation.formula1 or "",
                    "allow_blank": bool(validation.allowBlank),
                })
        return {
            "document_id": document_id,
            "sheet": sheet_name,
            "revision": int(record.get("revision") or 1),
            "max_row": max_row,
            "max_column": max_column,
            "columns": [
                {
                    "index": index,
                    "letter": get_column_letter(index),
                    "width": float(sheet.column_dimensions[get_column_letter(index)].width or 10),
                }
                for index in range(1, max_column + 1)
            ],
            "rows": [
                {"index": index, "height": float(sheet.row_dimensions[index].height or 20)}
                for index in range(1, max_row + 1)
            ],
            "cells": cells,
            "merges": [str(cell_range) for cell_range in sheet.merged_cells.ranges],
            "validations": validations,
        }

    def update_workbook_cells(
        self,
        project: dict[str, Any],
        document_id: str,
        sheet_name: str,
        cells: list[dict[str, Any]],
        expected_revision: int,
        actor: str,
    ) -> dict[str, Any]:
        index, record = self._get_record(project, document_id)
        if record.get("kind") != "workbook":
            raise DocumentWorkspaceError("当前文档不是表格文档")
        current_revision = int(record.get("revision") or 1)
        if expected_revision != current_revision:
            raise DocumentVersionConflict(f"表格已更新，当前版本为 v{current_revision}")
        path = self._content_path(project, record)
        workbook = load_workbook(path, read_only=False, data_only=False)
        if sheet_name not in workbook.sheetnames:
            raise DocumentWorkspaceError("工作表不存在")
        sheet = workbook[sheet_name]
        root = self._document_root(project, document_id)
        backup = root / "versions" / f"v{current_revision:04d}-before-{_safe_name(actor)}.xlsx"
        if not backup.exists():
            shutil.copy2(path, backup)
        for patch in cells:
            coordinate = str(patch.get("coordinate") or "").upper()
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,4}", coordinate):
                raise DocumentWorkspaceError(f"非法单元格坐标：{coordinate}")
            sheet[coordinate] = patch.get("value")
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        temporary = path.with_suffix(".xlsx.tmp")
        workbook.save(temporary)
        temporary.replace(path)
        next_revision = current_revision + 1
        record["revision"] = next_revision
        record["updated_at"] = _now()
        record.setdefault("metadata", {})["last_actor"] = actor
        record["metadata"]["recalculation"] = self._recalculate_workbook(path)
        self._write_index(project, index)
        return {
            "document": self._public_record(project, record),
            "revision": next_revision,
            "sheet": sheet_name,
            "updated_cells": len(cells),
            "recalculation": record["metadata"]["recalculation"],
        }

    def _recalculate_workbook(self, path: Path) -> dict[str, Any]:
        soffice = Path("/opt/homebrew/bin/soffice")
        if not soffice.exists():
            return {"status": "deferred", "message": "LibreOffice 不可用，公式将在 Excel 打开时重算"}
        with tempfile.TemporaryDirectory(prefix="openclaw-xlsx-", dir="/private/tmp") as directory:
            result = subprocess.run(
                [str(soffice), "--headless", "--convert-to", "xlsx", "--outdir", directory, str(path)],
                capture_output=True,
                text=True,
                timeout=90,
            )
            output = Path(directory) / path.name
            if result.returncode == 0 and output.exists() and output.stat().st_size:
                shutil.copy2(output, path)
                return {"status": "completed", "message": "公式已通过 LibreOffice 重算"}
            return {"status": "deferred", "message": (result.stderr or result.stdout or "公式重算失败")[-300:]}

    def rename_workbook_sheet(
        self,
        project: dict[str, Any],
        document_id: str,
        sheet_name: str,
        new_name: str,
        expected_revision: int,
    ) -> dict[str, Any]:
        _, record = self._get_record(project, document_id)
        if int(record.get("revision") or 1) != expected_revision:
            raise DocumentVersionConflict("表格版本已变化，请刷新后重试")
        path = self._content_path(project, record)
        workbook = load_workbook(path)
        if sheet_name not in workbook.sheetnames:
            raise DocumentWorkspaceError("工作表不存在")
        new_name = new_name.strip()
        if not new_name or len(new_name) > 31 or any(token in new_name for token in "[]:*?/\\"):
            raise DocumentWorkspaceError("工作表名称无效")
        if new_name in workbook.sheetnames and new_name != sheet_name:
            raise DocumentWorkspaceError("工作表名称已存在")
        workbook[sheet_name].title = new_name
        workbook.save(path)
        self._update_record(project, document_id, {"revision": expected_revision + 1, "updated_at": _now()})
        return self.workbook_metadata(project, document_id)

    def replace_content(self, project: dict[str, Any], document_id: str, data: bytes, filename: str, actor: str) -> dict[str, Any]:
        index, record = self._get_record(project, document_id)
        kind = record.get("kind")
        expected_suffix = ".xlsx" if kind == "workbook" else ".pptx" if kind == "presentation" else ""
        if not expected_suffix:
            raise DocumentWorkspaceError("正文文档不支持文件替换")
        if len(data) > MAX_UPLOAD_BYTES:
            raise DocumentWorkspaceError("文件超过 100 MB 上限")
        if Path(filename).suffix.lower() != expected_suffix:
            raise DocumentWorkspaceError(f"只支持 {expected_suffix} 文件")
        self._validate_office_zip(data, kind)
        path = self._content_path(project, record)
        root = self._document_root(project, document_id)
        path.parent.mkdir(parents=True, exist_ok=True)
        (root / "versions").mkdir(parents=True, exist_ok=True)
        current_revision = int(record.get("revision") or 1)
        if path.exists():
            shutil.copy2(path, root / "versions" / f"v{current_revision:04d}-before-{_safe_name(actor)}{expected_suffix}")
        temporary = path.with_suffix(expected_suffix + ".tmp")
        temporary.write_bytes(data)
        temporary.replace(path)
        if not (root / "source" / f"original{expected_suffix}").exists():
            (root / "source").mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, root / "source" / f"original{expected_suffix}")
        record["revision"] = current_revision + 1 if current_revision else 1
        record["source_path"] = filename
        record["source_checksum"] = hashlib.sha256(data).hexdigest()
        record["updated_at"] = _now()
        record.setdefault("metadata", {})["last_actor"] = actor
        if kind == "presentation":
            record["metadata"]["render"] = self._render_presentation(project, record)
            if isinstance(record["metadata"].get("structure_binding"), dict):
                record["metadata"]["structure_binding"]["status"] = "stale"
        elif kind == "workbook":
            record["metadata"]["recalculation"] = self._recalculate_workbook(path)
        self._write_index(project, index)
        return self._public_record(project, record)

    def _validate_office_zip(self, data: bytes, kind: str) -> None:
        with tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024) as handle:
            handle.write(data)
            handle.seek(0)
            try:
                with ZipFile(handle) as archive:
                    marker = "xl/workbook.xml" if kind == "workbook" else "ppt/presentation.xml"
                    names = archive.namelist()
                    if marker not in names:
                        raise DocumentWorkspaceError("Office 文件结构无效")
                    total = sum(item.file_size for item in archive.infolist())
                    if total > MAX_UPLOAD_BYTES * 10:
                        raise DocumentWorkspaceError("Office 文件解压规模异常")
            except BadZipFile as exc:
                raise DocumentWorkspaceError("Office 文件已损坏") from exc

    def _presentation_slide_count(self, path: Path) -> int:
        with ZipFile(path) as archive:
            return sum(1 for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name))

    def _presentation_pdf_path(self, project: dict[str, Any], record: dict[str, Any]) -> Path:
        return self._document_root(project, record["id"]) / "rendered" / "presentation.pdf"

    def _render_presentation(self, project: dict[str, Any], record: dict[str, Any]) -> dict[str, Any]:
        path = self._content_path(project, record)
        rendered = self._presentation_pdf_path(project, record)
        rendered.parent.mkdir(parents=True, exist_ok=True)
        soffice = Path("/opt/homebrew/bin/soffice")
        if not soffice.exists():
            return {"status": "failed", "message": "LibreOffice 不可用"}
        result = subprocess.run(
            [str(soffice), "--headless", "--convert-to", "pdf", "--outdir", str(rendered.parent), str(path)],
            capture_output=True,
            text=True,
            timeout=180,
        )
        generated = rendered.parent / f"{path.stem}.pdf"
        if generated.exists() and generated != rendered:
            generated.replace(rendered)
        if result.returncode != 0 or not rendered.exists() or not rendered.stat().st_size:
            return {"status": "failed", "message": (result.stderr or result.stdout or "PPT 渲染失败")[-500:]}
        return {
            "status": "completed",
            "slide_count": self._presentation_slide_count(path),
            "size_bytes": rendered.stat().st_size,
        }

    def source_file(self, project: dict[str, Any], document_id: str) -> Path:
        _, record = self._get_record(project, document_id)
        if record.get("kind") not in {"workbook", "presentation"}:
            raise DocumentWorkspaceError("当前文档没有可下载的 Office 源文件")
        path = self._content_path(project, record)
        if not path.exists():
            raise DocumentWorkspaceError("尚未上传源文件")
        return path

    def presentation_preview(self, project: dict[str, Any], document_id: str) -> Path:
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "presentation":
            raise DocumentWorkspaceError("当前文档不是演示文档")
        preview = self._presentation_pdf_path(project, record)
        if not preview.exists():
            result = self._render_presentation(project, record)
            if result.get("status") != "completed":
                raise DocumentWorkspaceError(str(result.get("message") or "PPT 预览生成失败"))
        return preview

    def structure_binding(
        self, project: dict[str, Any], document_id: str
    ) -> dict[str, Any]:
        _, record = self._get_record(project, document_id)
        return self._structure_binding(project, record)

    def set_structure_binding(
        self,
        project: dict[str, Any],
        document_id: str,
        binding: dict[str, Any],
        manifest: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        index, record = self._get_record(project, document_id)
        normalised = self._normalise_structure_binding(binding)
        source_document_id = normalised["source_document_id"]
        if not source_document_id or source_document_id == document_id:
            raise DocumentWorkspaceError("结构绑定缺少有效源文档")
        source = next(
            (
                row
                for row in index["documents"]
                if str(row.get("id") or "") == source_document_id
            ),
            None,
        )
        if source is None:
            raise DocumentWorkspaceError("结构绑定源文档不存在")
        normalised["source_sha256"] = (
            normalised["source_sha256"] or self._record_checksum(project, source)
        )
        normalised["source_version"] = (
            normalised["source_version"] or self._record_structure_version(source)
        )

        presentation_summary: dict[str, Any] = {}
        if manifest is not None:
            if record.get("kind") != "presentation":
                raise DocumentWorkspaceError("只有演示文档支持逐页联动清单")
            validated_manifest = self._validate_presentation_manifest(manifest)
            slides = validated_manifest.get("slides") or []
            unmapped = [
                int(item["slide"])
                for item in slides
                if not item.get("thesis_sections")
            ]
            normalised["mapped_items"] = len(slides) - len(unmapped)
            normalised["unmapped_items"] = list(
                dict.fromkeys([*normalised["unmapped_items"], *unmapped])
            )
            presentation_output = (
                validated_manifest.get("authority", {}).get(
                    "presentation_output", {}
                )
                if isinstance(validated_manifest.get("authority"), dict)
                else {}
            )
            presentation_summary = {
                "slide_count": len(slides),
                "speaker_notes_count": int(
                    presentation_output.get("notes_count") or 0
                ),
                "main_slide_count": int(
                    presentation_output.get("main_slide_count") or 0
                ),
                "appendix_slide_count": int(
                    presentation_output.get("appendix_slide_count") or 0
                ),
                "mapping_contract_version": str(
                    validated_manifest.get("contract_version") or ""
                ),
            }
            validated_manifest["structure_binding"] = normalised
            self._write_json_atomic(
                self._presentation_manifest_path(project, document_id),
                validated_manifest,
            )
        elif record.get("kind") == "presentation":
            existing_manifest = self._read_presentation_manifest(project, document_id)
            if existing_manifest:
                existing_manifest["structure_binding"] = normalised
                self._write_json_atomic(
                    self._presentation_manifest_path(project, document_id),
                    existing_manifest,
                )

        record.setdefault("metadata", {})["structure_binding"] = normalised
        if presentation_summary:
            record["metadata"].update(presentation_summary)
        record["updated_at"] = _now()
        self._write_index(project, index)
        evaluated = self._structure_binding(project, record)
        if evaluated != normalised:
            index, record = self._get_record(project, document_id)
            record.setdefault("metadata", {})["structure_binding"] = evaluated
            self._write_index(project, index)
            manifest_path = self._presentation_manifest_path(project, document_id)
            if manifest_path.exists():
                current_manifest = self._read_presentation_manifest(project, document_id)
                current_manifest["structure_binding"] = evaluated
                self._write_json_atomic(manifest_path, current_manifest)
        return evaluated

    def presentation_manifest(
        self, project: dict[str, Any], document_id: str
    ) -> dict[str, Any]:
        _, record = self._get_record(project, document_id)
        if record.get("kind") != "presentation":
            raise DocumentWorkspaceError("当前文档不是演示文档")
        manifest = self._read_presentation_manifest(project, document_id)
        if not manifest:
            manifest = {
                "schema": "openclaw.presentation-manifest",
                "contract_version": "",
                "narrative_sections": [],
                "slides": [],
            }
        manifest["structure_binding"] = self._structure_binding(project, record)
        manifest["document"] = self._public_record(project, record)
        return manifest

    def presentation_slide_proposal(
        self,
        project: dict[str, Any],
        document_id: str,
        slide: int,
        draft: dict[str, Any],
        instruction: str,
        agent_id: str = "presentation-editor",
    ) -> dict[str, Any]:
        manifest = self.presentation_manifest(project, document_id)
        slides = manifest.get("slides") or []
        if slide < 1:
            raise DocumentWorkspaceError("PPT 页码无效")
        current = next(
            (row for row in slides if int(row.get("slide") or 0) == slide),
            {"slide": slide},
        )
        title = self._compact_presentation_title(
            str(draft.get("title") or current.get("title") or "")
        )
        claim = str(draft.get("claim") or current.get("claim") or "").strip()
        evidence_level = str(
            draft.get("evidence_level") or current.get("evidence_level") or "C级"
        ).strip()
        if evidence_level not in {"A级", "B级", "C级", "D级"}:
            evidence_level = "C级"
        instruction_text = instruction.strip()
        if not claim:
            claim = "本页需要补充清晰的核心主张。"
        if "证据" not in claim and "边界" not in claim and evidence_level != "A级":
            claim = f"{claim.rstrip('。')}（证据等级：{evidence_level}，不超出已登记材料作强结论。）"

        notes = str(draft.get("notes") or current.get("notes") or "").strip()
        if "压缩" in instruction_text:
            notes = self._trim_presentation_notes(notes)
        if not notes:
            notes = f"本页先说明“{title}”的核心问题，再解释其与正文结构的关系。"
        boundary = (
            "证据口径：可作为正式实验结论陈述。"
            if evidence_level == "A级"
            else "证据口径：当前不是A级证据，只能作为阶段性依据或方法说明。"
        )
        if boundary not in notes:
            notes = f"{notes}\n\n讲解重点：围绕“{claim.rstrip('。')}”展开，避免新增未核验事实。\n{boundary}"

        evidence_ids = draft.get("evidence_ids")
        if not isinstance(evidence_ids, list):
            evidence_ids = current.get("evidence_ids") if isinstance(current, dict) else []
        evidence_ids = evidence_ids or []
        evidence_ids = [str(item).strip() for item in evidence_ids if str(item).strip()]
        patch = {
            "title": title or "未命名页面",
            "claim": claim,
            "notes": notes,
            "evidence_level": evidence_level,
            "evidence_ids": evidence_ids,
            "appendix": bool(draft.get("appendix", current.get("appendix", False))),
        }
        proposal_seed = json.dumps(
            {
                "project": project.get("id"),
                "document": document_id,
                "slide": slide,
                "instruction": instruction_text,
                "patch": patch,
            },
            ensure_ascii=False,
            sort_keys=True,
        )
        return {
            "schema": "openclaw.presentation-slide-proposal.v1",
            "id": f"ppt-proposal-{hashlib.sha1(proposal_seed.encode()).hexdigest()[:12]}",
            "status": "succeeded",
            "agent_id": agent_id,
            "slide": slide,
            "title": "AI 页面微调建议",
            "summary": "已根据当前页标题、主张、讲解稿和证据等级生成结构化修改建议。",
            "rationale": "P1 阶段只返回结构化 manifest patch，不直接修改 PPTX。",
            "patch": patch,
        }

    def submit_presentation_slide_job(
        self,
        project: dict[str, Any],
        document_id: str,
        slide: int,
        draft: dict[str, Any],
        instruction: str,
        agent_id: str = "presentation-editor",
        client_request_id: str = "",
    ) -> dict[str, Any]:
        seed = json.dumps(
            {
                "project": project.get("id"),
                "document": document_id,
                "slide": slide,
                "instruction": instruction,
                "draft": draft,
                "client_request_id": client_request_id,
            },
            ensure_ascii=False,
            sort_keys=True,
        )
        job_id = f"ppt-job-{hashlib.sha1(seed.encode()).hexdigest()[:12]}"
        existing = self._presentation_slide_jobs.get(job_id)
        if existing:
            return copy.deepcopy(existing)

        now = _now()
        job = {
            "id": job_id,
            "status": "running",
            "project_id": str(project.get("id") or ""),
            "document_id": document_id,
            "slide": slide,
            "agent_id": agent_id,
            "instruction": instruction,
            "created_at": now,
            "updated_at": now,
            "proposal": None,
            "error": "",
        }
        self._presentation_slide_jobs[job_id] = job
        try:
            proposal = self.presentation_slide_proposal(
                project, document_id, slide, draft, instruction, agent_id
            )
            job.update({
                "status": "succeeded",
                "proposal": proposal,
                "updated_at": _now(),
            })
        except Exception as exc:
            job.update({
                "status": "failed",
                "error": str(exc),
                "updated_at": _now(),
            })
        return copy.deepcopy(job)

    def get_presentation_slide_job(
        self, project: dict[str, Any], document_id: str, job_id: str
    ) -> dict[str, Any]:
        job = self._presentation_slide_jobs.get(job_id)
        if (
            not job
            or job.get("project_id") != str(project.get("id") or "")
            or job.get("document_id") != document_id
        ):
            raise DocumentWorkspaceError("PPT 页面 AI 任务不存在")
        return copy.deepcopy(job)

    def _compact_presentation_title(self, value: str) -> str:
        title = re.split(r"[，,。；;：:]", value.replace("\n", " "))[0].strip()
        return title[:24] if title else ""

    def _trim_presentation_notes(self, value: str) -> str:
        lines = [line.strip() for line in value.splitlines() if line.strip()]
        return "\n".join(lines[:4])

    def versions(self, project: dict[str, Any], document_id: str) -> list[dict[str, Any]]:
        _, record = self._get_record(project, document_id)
        if record.get("kind") == "rich_text":
            return document_workspace_service.versions(self._rich_project(project, record))
        root = self._document_root(project, document_id)
        rows = []
        for path in sorted((root / "versions").glob("*"), key=lambda item: item.stat().st_mtime, reverse=True):
            stat = path.stat()
            rows.append({
                "name": path.name,
                "size_bytes": stat.st_size,
                "created_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            })
        current = self._content_path(project, record)
        if current.exists():
            stat = current.stat()
            rows.insert(0, {
                "name": f"当前版本 v{record.get('revision') or 1}",
                "size_bytes": stat.st_size,
                "created_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                "current": True,
            })
        return rows

    def restore_version(
        self,
        project: dict[str, Any],
        document_id: str,
        version_name: str,
        actor: str,
    ) -> dict[str, Any]:
        index, record = self._get_record(project, document_id)
        if record.get("kind") not in {"workbook", "presentation"}:
            raise DocumentWorkspaceError("正文文档暂不支持通过此接口恢复版本")
        if Path(version_name).name != version_name:
            raise DocumentWorkspaceError("非法版本名称")
        root = self._document_root(project, document_id)
        version_path = (root / "versions" / version_name).resolve()
        try:
            version_path.relative_to((root / "versions").resolve())
        except ValueError as exc:
            raise DocumentWorkspaceError("非法版本路径") from exc
        if not version_path.is_file():
            raise DocumentWorkspaceError("版本文件不存在")
        current = self._content_path(project, record)
        current_revision = int(record.get("revision") or 1)
        suffix = current.suffix
        backup = root / "versions" / f"v{current_revision:04d}-before-restore-{_safe_name(actor)}{suffix}"
        if current.exists() and not backup.exists():
            shutil.copy2(current, backup)
        temporary = current.with_suffix(suffix + ".restore")
        shutil.copy2(version_path, temporary)
        temporary.replace(current)
        record["revision"] = current_revision + 1
        record["updated_at"] = _now()
        record.setdefault("metadata", {})["restored_from"] = version_name
        record["metadata"]["last_actor"] = actor
        if record["kind"] == "presentation":
            record["metadata"]["render"] = self._render_presentation(project, record)
            if isinstance(record["metadata"].get("structure_binding"), dict):
                record["metadata"]["structure_binding"]["status"] = "stale"
        else:
            record["metadata"]["recalculation"] = self._recalculate_workbook(current)
        self._write_index(project, index)
        return self._public_record(project, record)


multi_document_service = MultiDocumentService()
