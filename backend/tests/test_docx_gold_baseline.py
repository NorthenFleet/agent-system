import base64
import hashlib
import io
import json
import subprocess
import sys
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches

from services.docx_gold_baseline import audit_docx, audit_manifest


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZKMcAAAAASUVORK5CYII="
)


def _gold_fixture(tmp_path):
    document = Document()
    document.add_heading("第一章 金标", level=1)
    body = document.add_paragraph("人工与AI围绕同一份文档协作。")
    body._p.append(parse_xml(f'<m:oMath {nsdecls("m")}><m:r><m:t>x+y</m:t></m:r></m:oMath>'))
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), "7")
    bookmark_start.set(qn("w:name"), "gold-anchor")
    body._p.insert(0, bookmark_start)
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), "7")
    body._p.append(bookmark_end)

    field = document.add_paragraph()
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field.add_run()._r.append(instruction)

    complex_field = document.add_paragraph()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    complex_field.add_run()._r.append(begin)
    for fragment in (" REF ", "gold-", "anchor "):
        part = OxmlElement("w:instrText")
        part.set(qn("xml:space"), "preserve")
        part.text = fragment
        complex_field.add_run()._r.append(part)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    complex_field.add_run()._r.append(separate)
    complex_field.add_run("1")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    complex_field.add_run()._r.append(end)

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并"
    table.cell(1, 0).text = "指标"
    table.cell(1, 1).text = "结果"
    image = tmp_path / "pixel.png"
    image.write_bytes(PNG_1X1)
    document.add_picture(str(image), width=Inches(1))
    document.sections[0].header.paragraphs[0].text = "金标页眉"
    stream = io.BytesIO()
    document.save(stream)
    path = tmp_path / "gold.docx"
    path.write_bytes(stream.getvalue())
    return path


def _mutate_zip_part(path, tmp_path, part_name, old, new):
    output = tmp_path / f"mutated-{path.name}"
    with ZipFile(path) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == part_name:
                data = data.replace(old, new)
            target.writestr(item, data)
    return output


def test_docx_gold_inventory_captures_native_ooxml_objects(tmp_path):
    path = _gold_fixture(tmp_path)
    expected = hashlib.sha256(path.read_bytes()).hexdigest()

    report = audit_docx(path, expected_sha256=expected)

    assert report["status"] == "passed"
    assert report["package"]["missing_required_parts"] == []
    assert report["metrics"]["heading_count"] == 1
    assert report["metrics"]["table_count"] == 1
    assert report["metrics"]["horizontal_merge_region_count"] == 1
    assert report["metrics"]["horizontal_merged_cell_extra_count"] == 1
    assert report["metrics"]["media_part_count"] == 1
    assert report["metrics"]["math_object_count"] == 1
    assert report["metrics"]["bookmark_count"] == 1
    assert report["field_type_counts"]["PAGE"] == 1
    assert report["field_type_counts"]["REF"] == 1
    assert "OTHER" not in report["field_type_counts"]
    assert report["bookmark_names"] == ["gold-anchor"]
    assert len(report["text_sha256"]) == 64
    assert len(report["structure_sha256"]) == 64
    assert len(report["native_object_sha256"]["math"]) == 64
    assert len(report["native_object_sha256"]["headers"]) == 64


def test_docx_gold_inventory_fails_closed_on_hash_mismatch(tmp_path):
    path = _gold_fixture(tmp_path)

    report = audit_docx(path, expected_sha256="0" * 64)

    assert report["status"] == "failed"
    assert report["errors"] == ["sha256_mismatch"]


def test_docx_gold_inventory_requires_a_valid_expected_hash(tmp_path):
    path = _gold_fixture(tmp_path)

    missing = audit_docx(path)
    malformed = audit_docx(path, expected_sha256="not-a-sha")

    assert missing["status"] == "failed"
    assert "expected_sha256_required" in missing["errors"]
    assert malformed["status"] == "failed"
    assert "expected_sha256_invalid" in malformed["errors"]


def test_native_object_hashes_detect_formula_and_header_damage(tmp_path):
    path = _gold_fixture(tmp_path)
    expected = hashlib.sha256(path.read_bytes()).hexdigest()
    original = audit_docx(path, expected_sha256=expected)

    formula_path = _mutate_zip_part(path, tmp_path, "word/document.xml", b"x+y", b"x-y")
    formula = audit_docx(
        formula_path,
        expected_sha256=hashlib.sha256(formula_path.read_bytes()).hexdigest(),
    )
    header_path = _mutate_zip_part(
        path,
        tmp_path,
        "word/header1.xml",
        "金标页眉".encode(),
        "损坏页眉".encode(),
    )
    header = audit_docx(
        header_path,
        expected_sha256=hashlib.sha256(header_path.read_bytes()).hexdigest(),
    )

    assert formula["text_sha256"] == original["text_sha256"]
    assert formula["native_object_sha256"]["math"] != original["native_object_sha256"]["math"]
    assert formula["structure_sha256"] != original["structure_sha256"]
    assert header["native_object_sha256"]["headers"] != original["native_object_sha256"]["headers"]
    assert header["structure_sha256"] != original["structure_sha256"]


def test_manifest_report_is_deterministic_and_preserves_authority_roles(tmp_path):
    path = _gold_fixture(tmp_path)
    expected = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(
        json.dumps(
            {
                "schema": "openclaw.word-fidelity-gold-set.v1",
                "documents": [
                    {
                        "id": "fixture",
                        "title": "金标",
                        "path": str(path),
                        "expected_sha256": expected,
                        "authority_role": "layout_reference",
                        "authority_evidence": ["test"],
                    }
                ],
                "acceptance_profiles": {"no_op_roundtrip": {"required": True}},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    first = audit_manifest(manifest_path)
    second = audit_manifest(manifest_path)
    manifest_link = tmp_path / "manifest-link.json"
    manifest_link.symlink_to(manifest_path)
    via_link = audit_manifest(manifest_link)

    assert first == second
    assert first["status"] == "passed"
    assert first["documents"][0]["authority_role"] == "layout_reference"
    assert first["acceptance_profiles"]["no_op_roundtrip"]["required"] is True
    assert len(first["report_sha256"]) == 64
    assert via_link["manifest_path"] != first["manifest_path"]
    assert via_link["report_sha256"] == first["report_sha256"]


def test_verify_failure_does_not_replace_existing_baseline(tmp_path):
    path = _gold_fixture(tmp_path)
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(
        json.dumps(
            {
                "documents": [
                    {
                        "id": "fixture",
                        "path": str(path),
                        "expected_sha256": "0" * 64,
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    output = tmp_path / "baseline.json"
    output.write_text("known-good\n", encoding="utf-8")
    script = __file__.replace("tests/test_docx_gold_baseline.py", "scripts/audit_docx_gold_baseline.py")

    result = subprocess.run(
        [sys.executable, script, "--manifest", str(manifest_path), "--output", str(output), "--verify"],
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 1
    assert output.read_text(encoding="utf-8") == "known-good\n"
    assert "baseline not replaced" in result.stderr
