import base64
import io
import shutil

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches

from services.pandoc_document_importer import PandocDocxImporter


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZKMcAAAAASUVORK5CYII="
)


def omath(value: str):
    return parse_xml(
        f'<m:oMath {nsdecls("m")}><m:r><m:t>{value}</m:t></m:r></m:oMath>'
    )


@pytest.mark.skipif(not shutil.which("pandoc"), reason="Pandoc is required for fidelity imports")
def test_docx_ast_import_preserves_math_media_tables_and_unicode(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("字符：α β ≤ ≥ → 舰艇")
    paragraph._p.append(omath("x+y"))
    image = tmp_path / "pixel.png"
    image.write_bytes(PNG_1X1)
    document.add_picture(str(image), width=Inches(1))
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "合并"
    table.cell(0, 1)._tc.get_or_add_tcPr()
    table.cell(0, 1).paragraphs[0]._p.append(omath("a=b"))
    table.cell(0, 2).text = "（1.1）"
    table.cell(1, 1).text = "数据"
    table.cell(1, 2).text = "说明"
    stream = io.BytesIO()
    document.save(stream)

    result = PandocDocxImporter().convert(stream.getvalue(), "高保真测试")
    metrics = result["stats"]

    assert metrics["fidelity_status"] == "passed"
    assert metrics["replacement_character_count"] == 0
    assert metrics["image_count"] == 1
    assert metrics["table_count"] == 1
    assert metrics["math_block_count"] + metrics["math_inline_count"] == 2
    assert len(result["assets"]) == 1
    assert "α β ≤ ≥ → 舰艇" in result["markdown"]
    table_node = next(node for node in result["document"]["content"] if node["type"] == "table")
    assert any(
        cell["attrs"]["rowspan"] == 2
        for row in table_node["content"]
        for cell in row["content"]
    )
