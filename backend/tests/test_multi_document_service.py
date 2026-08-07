import hashlib
import io
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from docx import Document

from services.document_workspace_service import DocumentVersionConflict, document_workspace_service
from services.multi_document_service import MultiDocumentService


@pytest.fixture
def service(tmp_path, monkeypatch):
    monkeypatch.setattr(document_workspace_service, "vault", tmp_path)
    monkeypatch.setattr(document_workspace_service, "_find_source_word", lambda project: None)
    source = tmp_path / "10-成果库-Outputs" / "test" / "body.md"
    source.parent.mkdir(parents=True)
    source.write_text(
        "# 文档说明\n\n测试正文。\n\n# 第一章 测试\n\n正文内容。\n\n# 参考文献\n\n待补充。\n",
        encoding="utf-8",
    )
    return MultiDocumentService(), source


def project(source: Path):
    return {
        "id": "proj-test",
        "name": "测试·多文档项目",
        "enabled_modules": ["writing"],
        "document_spec": {
            "working_markdown": {"path": str(source), "chapter_count": 1},
            "expected_chapters": 1,
        },
    }


def presentation_bytes(marker: str = "v1", slide_count: int = 2) -> bytes:
    buffer = io.BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("ppt/presentation.xml", f"<presentation>{marker}</presentation>")
        for index in range(1, slide_count + 1):
            archive.writestr(f"ppt/slides/slide{index}.xml", f"<slide>{marker}-{index}</slide>")
    return buffer.getvalue()


def docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("课程导入材料", level=1)
    document.add_paragraph("这是导入后的教学材料正文。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "讲次"
    table.cell(0, 1).text = "学时"
    table.cell(1, 0).text = "第1讲"
    table.cell(1, 1).text = "2"
    document.save(buffer)
    return buffer.getvalue()


def test_legacy_workspace_becomes_primary_rich_document(service):
    multi, source = service
    result = multi.list_documents(project(source))
    assert result["summary"]["total"] == 1
    primary = result["documents"][0]
    assert primary["kind"] == "rich_text"
    assert primary["is_primary"] is True
    workspace = multi.rich_workspace(project(source), primary["id"])
    assert workspace["stats"]["chapter_count"] == 1


def test_workbook_roundtrip_keeps_formula_and_versions(service, tmp_path):
    multi, source = service
    workbook_source = tmp_path / "10-成果库-Outputs" / "test" / "table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "裁决"
    sheet["A1"] = 2
    sheet["B1"] = 3
    sheet["C1"] = "=A1+B1"
    workbook.save(workbook_source)

    document = multi.create_document(project(source), "裁决表", "workbook", source_path=str(workbook_source))
    metadata = multi.workbook_metadata(project(source), document["id"])
    assert metadata["sheet_count"] == 1
    assert metadata["formula_count"] == 1

    sheet_data = multi.workbook_sheet(project(source), document["id"], "裁决")
    formula = next(cell for cell in sheet_data["cells"] if cell["coordinate"] == "C1")
    assert formula["formula"] == "=A1+B1"

    updated = multi.update_workbook_cells(
        project(source),
        document["id"],
        "裁决",
        [{"coordinate": "A1", "value": 4}],
        sheet_data["revision"],
        "tester",
    )
    assert updated["revision"] == sheet_data["revision"] + 1
    version_name = next(
        row["name"]
        for row in multi.versions(project(source), document["id"])
        if row["name"].startswith("v0001-before")
    )
    restored = multi.restore_version(project(source), document["id"], version_name, "tester")
    assert restored["revision"] == updated["revision"] + 1
    with pytest.raises(DocumentVersionConflict):
        multi.update_workbook_cells(
            project(source),
            document["id"],
            "裁决",
            [{"coordinate": "A1", "value": 5}],
            sheet_data["revision"],
            "tester",
        )


def test_document_archive_restore_and_delete(service):
    multi, source = service
    document = multi.create_document(project(source), "临时正文", "rich_text")
    archived = multi.update_document(project(source), document["id"], {"status": "archived"})
    assert archived["status"] == "archived"
    restored = multi.update_document(project(source), document["id"], {"status": "active"})
    assert restored["status"] == "active"
    multi.update_document(project(source), document["id"], {"status": "archived"})
    multi.delete_document(project(source), document["id"])
    assert all(row["id"] != document["id"] for row in multi.list_documents(project(source))["documents"])


def test_output_metadata_defaults_and_internal_workbook(service, tmp_path):
    multi, source = service
    formal = multi.create_document(project(source), "裁决规则", "rich_text")
    assert formal["is_output_product"] is True
    assert formal["output_format"] == "docx"
    assert formal["publication_status"] == "draft"

    workbook_source = tmp_path / "10-成果库-Outputs" / "test" / "data.xlsx"
    workbook_source.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.save(workbook_source)
    internal = multi.create_document(
        project(source),
        "裁决数据源",
        "workbook",
        source_path=str(workbook_source),
    )
    assert internal["is_output_product"] is False
    assert internal["publication_status"] == "internal"

    updated = multi.update_document(
        project(source),
        formal["id"],
        {
            "data_source_ids": [internal["id"]],
            "print_profile": "wargame_a4",
            "publication_status": "approved",
            "rules_version": "0522-v1",
            "data_version": "0522-data-v1",
        },
    )
    assert updated["data_source_ids"] == [internal["id"]]
    assert updated["print_profile"] == "wargame_a4"
    assert updated["publication_status"] == "approved"


def test_course_product_metadata_and_docx_fidelity_import(service):
    multi, source = service
    original_docx = docx_bytes()
    imported = multi.import_docx(
        project(source),
        "课程教学材料",
        original_docx,
        "课程教学材料.docx",
        product_type="lesson_plan",
        course_unit_ids=["L01", "L01"],
        quality_profile="lesson_plan_v1",
        required_for_release=True,
        source_refs=[{
            "ref_type": "project_document",
            "project_id": "proj-manual",
            "document_id": "doc-rules",
            "relation": "authoritative_rule",
            "required": True,
            "version": "R1.2/D1.2",
        }],
    )

    assert imported["product_type"] == "lesson_plan"
    assert imported["course_unit_ids"] == ["L01"]
    assert imported["required_for_release"] is True
    assert imported["source_refs"][0]["version"] == "R1.2/D1.2"
    assert imported["import_summary"]["heading_count"] == 1
    assert imported["import_summary"]["table_count"] == 1
    fulltext = multi.rich_call(project(source), imported["id"], "fulltext")["content"]
    assert "# 课程导入材料" in fulltext
    assert "| 讲次 | 学时 |" in fulltext
    source_word = Path(
        multi.rich_call(project(source), imported["id"], "ensure_workspace")["source_word"]
    )
    assert source_word.name == "original.docx"
    assert source_word.read_bytes() == original_docx


def test_secondary_document_uses_its_own_expected_chapter_count(service):
    multi, source = service
    document = multi.create_document(
        project(source),
        "裁决规则",
        "rich_text",
        outline=["第一章 裁决总则", "第二章 状态体系"],
    )

    assert document["expected_chapters"] == 2
    workspace = multi.rich_workspace(project(source), document["id"])
    assert workspace["stats"]["chapter_count"] == 2
    assert workspace["quality"]["blockers"] == 0
    assert workspace["quality"]["score"] == 100


def test_upload_rich_text_asset_writes_document_asset(service):
    multi, source = service
    document = multi.create_document(project(source), "配图正文", "rich_text")
    uploaded = multi.upload_rich_text_asset(
        project(source),
        document["id"],
        b"fake-png",
        "../体系结构.png",
        "image/png",
        actor="tester",
    )

    assert uploaded["path"].startswith("assets/")
    assert uploaded["filename"].endswith(".png")
    assert ".." not in uploaded["path"]
    resolved = multi.rich_call(project(source), document["id"], "asset_path", uploaded["path"])
    assert resolved.read_bytes() == b"fake-png"

    updated = next(row for row in multi.list_documents(project(source))["documents"] if row["id"] == document["id"])
    uploads = updated["metadata"]["asset_uploads"]
    assert uploads[-1]["path"] == uploaded["path"]
    assert uploads[-1]["actor"] == "tester"


def test_upload_rich_text_asset_rejects_non_images(service):
    multi, source = service
    document = multi.create_document(project(source), "配图正文", "rich_text")

    with pytest.raises(Exception, match="只支持"):
        multi.upload_rich_text_asset(
            project(source),
            document["id"],
            b"not-image",
            "payload.txt",
            "text/plain",
        )


def test_export_package_contains_only_approved_output_docx(service, tmp_path, monkeypatch):
    multi, source = service
    primary = multi.list_documents(project(source))["documents"][0]
    multi.update_document(
        project(source),
        primary["id"],
        {"publication_status": "approved", "print_profile": "wargame_a4"},
    )
    adjudication = multi.create_document(
        project(source),
        "裁决规则",
        "rich_text",
        publication_status="approved",
        print_profile="wargame_a4",
    )
    workbook_source = tmp_path / "10-成果库-Outputs" / "test" / "internal.xlsx"
    workbook_source.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.save(workbook_source)
    multi.create_document(
        project(source),
        "裁决数据源",
        "workbook",
        source_path=str(workbook_source),
    )

    def fake_export(synthetic_project, output_format):
        target = tmp_path / f"{synthetic_project.get('_document_id')}.docx"
        target.write_bytes(f"{synthetic_project.get('_document_title')}:{output_format}".encode())
        return target

    monkeypatch.setattr(document_workspace_service, "export", fake_export)
    package = multi.export_package(project(source))
    with ZipFile(package) as archive:
        names = archive.namelist()
    assert len(names) == 2
    assert all(name.endswith(".docx") for name in names)
    assert all(name.startswith("测试·多文档项目——") for name in names)
    assert any("正文" in name for name in names)
    assert any(adjudication["title"] in name for name in names)
    assert all("裁决数据源" not in name for name in names)


def test_presentation_structure_binding_detects_stale_and_preserves_versions(
    service, tmp_path, monkeypatch
):
    multi, source = service
    monkeypatch.setattr(
        multi,
        "_render_presentation",
        lambda _project, _record: {"status": "completed", "slide_count": 2},
    )
    ppt_source = tmp_path / "10-成果库-Outputs" / "test" / "defense.pptx"
    ppt_source.write_bytes(presentation_bytes())
    thesis = multi.list_documents(project(source))["documents"][0]
    presentation = multi.create_document(
        project(source),
        "答辩PPT",
        "presentation",
        source_path=str(ppt_source),
        is_output_product=True,
        output_format="pptx",
        publication_status="draft",
        rules_version="结构v14",
        data_version="文档v3",
    )
    current_sha = hashlib.sha256(presentation_bytes()).hexdigest()
    manifest = {
        "schema": "openclaw.paper-ppt-system-mapping",
        "contract_version": "v14.1",
        "authority": {
            "presentation_output": {
                "sha256": current_sha,
                "slide_count": 2,
                "main_slide_count": 1,
                "appendix_slide_count": 1,
                "notes_count": 2,
            }
        },
        "slides": [
            {"slide": 1, "title": "正文", "thesis_sections": ["1.1"], "notes": "正文"},
            {"slide": 2, "title": "附录", "thesis_sections": ["附录"], "notes": "附录"},
        ],
    }
    binding = multi.set_structure_binding(
        project(source),
        presentation["id"],
        {
            "mode": "mapped",
            "source_document_id": thesis["id"],
            "source_version": "v1",
            "source_sha256": thesis["source_checksum"],
            "status": "aligned",
            "mapped_items": 2,
            "unmapped_items": [],
            "changed_sections": [],
        },
        manifest,
    )
    assert binding["status"] == "aligned"
    linked = multi.get_document(project(source), presentation["id"])
    assert linked["is_output_product"] is True
    assert linked["output_format"] == "pptx"
    assert linked["stats"]["slide_count"] == 2
    assert linked["stats"]["main_slide_count"] == 1
    assert linked["stats"]["appendix_slide_count"] == 1
    assert linked["stats"]["notes_count"] == 2

    replaced = multi.replace_content(
        project(source),
        presentation["id"],
        presentation_bytes("v2"),
        "defense-v2.pptx",
        "tester",
    )
    assert replaced["structure_binding"]["status"] == "stale"
    assert any(
        row["name"].startswith("v0001-before")
        for row in multi.versions(project(source), presentation["id"])
    )


def test_presentation_manifest_rejects_duplicate_and_reports_unmapped(
    service, tmp_path, monkeypatch
):
    multi, source = service
    monkeypatch.setattr(
        multi,
        "_render_presentation",
        lambda _project, _record: {"status": "completed", "slide_count": 2},
    )
    ppt_source = tmp_path / "10-成果库-Outputs" / "test" / "mapping.pptx"
    ppt_source.write_bytes(presentation_bytes())
    thesis = multi.list_documents(project(source))["documents"][0]
    presentation = multi.create_document(
        project(source), "联动稿", "presentation", source_path=str(ppt_source)
    )
    base_binding = {
        "mode": "mapped",
        "source_document_id": thesis["id"],
        "source_version": "v1",
        "source_sha256": thesis["source_checksum"],
        "status": "aligned",
        "mapped_items": 0,
        "unmapped_items": [],
        "changed_sections": [],
    }
    duplicate = {
        "authority": {"presentation_output": {"sha256": presentation["source_checksum"]}},
        "slides": [
            {"slide": 1, "thesis_sections": ["1.1"]},
            {"slide": 1, "thesis_sections": ["1.2"]},
        ],
    }
    with pytest.raises(Exception, match="重复页码"):
        multi.set_structure_binding(
            project(source), presentation["id"], base_binding, duplicate
        )

    missing_map = {
        "authority": {"presentation_output": {"sha256": presentation["source_checksum"]}},
        "slides": [
            {"slide": 1, "thesis_sections": ["1.1"]},
            {"slide": 2, "thesis_sections": []},
        ],
    }
    binding = multi.set_structure_binding(
        project(source), presentation["id"], base_binding, missing_map
    )
    assert binding["status"] == "diverged"
    assert binding["mapped_items"] == 1
    assert binding["unmapped_items"] == [2]


def test_presentation_slide_proposal_returns_structured_patch(
    service, tmp_path, monkeypatch
):
    multi, source = service
    monkeypatch.setattr(
        multi,
        "_render_presentation",
        lambda _project, _record: {"status": "completed", "slide_count": 1},
    )
    ppt_source = tmp_path / "10-成果库-Outputs" / "test" / "proposal.pptx"
    ppt_source.write_bytes(presentation_bytes(slide_count=1))
    thesis = multi.list_documents(project(source))["documents"][0]
    presentation = multi.create_document(
        project(source), "答辩稿", "presentation", source_path=str(ppt_source)
    )
    binding = {
        "mode": "mapped",
        "source_document_id": thesis["id"],
        "source_version": "v1",
        "source_sha256": thesis["source_checksum"],
        "status": "aligned",
        "mapped_items": 1,
        "unmapped_items": [],
        "changed_sections": [],
    }
    manifest = {
        "authority": {"presentation_output": {"sha256": presentation["source_checksum"]}},
        "slides": [{
            "slide": 1,
            "title": "理论基础，详细说明",
            "thesis_sections": ["2.1"],
            "claim": "指挥控制连接人的决心与模型算法。",
            "evidence_level": "C级",
            "evidence_ids": ["MODEL-2.1"],
            "notes": "",
            "appendix": False,
        }],
    }
    multi.set_structure_binding(project(source), presentation["id"], binding, manifest)

    proposal = multi.presentation_slide_proposal(
        project(source),
        presentation["id"],
        1,
        {
            "title": "理论基础，详细说明",
            "claim": "指挥控制连接人的决心与模型算法。",
            "evidence_level": "C级",
            "evidence_ids": ["MODEL-2.1"],
            "appendix": False,
        },
        "压缩标题并补充证据边界",
    )

    assert proposal["schema"] == "openclaw.presentation-slide-proposal.v1"
    assert proposal["patch"]["title"] == "理论基础"
    assert "证据等级：C级" in proposal["patch"]["claim"]
    assert "当前不是A级证据" in proposal["patch"]["notes"]
    assert proposal["patch"]["evidence_ids"] == ["MODEL-2.1"]


def test_presentation_slide_job_can_be_submitted_and_fetched(
    service, tmp_path, monkeypatch
):
    multi, source = service
    monkeypatch.setattr(
        multi,
        "_render_presentation",
        lambda _project, _record: {"status": "completed", "slide_count": 1},
    )
    ppt_source = tmp_path / "10-成果库-Outputs" / "test" / "job.pptx"
    ppt_source.write_bytes(presentation_bytes(slide_count=1))
    thesis = multi.list_documents(project(source))["documents"][0]
    presentation = multi.create_document(
        project(source), "答辩稿", "presentation", source_path=str(ppt_source)
    )
    binding = {
        "mode": "mapped",
        "source_document_id": thesis["id"],
        "source_version": "v1",
        "source_sha256": thesis["source_checksum"],
        "status": "aligned",
        "mapped_items": 1,
        "unmapped_items": [],
        "changed_sections": [],
    }
    manifest = {
        "authority": {"presentation_output": {"sha256": presentation["source_checksum"]}},
        "slides": [{"slide": 1, "title": "理论基础", "thesis_sections": ["2.1"]}],
    }
    multi.set_structure_binding(project(source), presentation["id"], binding, manifest)

    job = multi.submit_presentation_slide_job(
        project(source),
        presentation["id"],
        1,
        {"title": "理论基础", "claim": "需要说明证据边界。"},
        "强化表达",
        client_request_id="client-1",
    )
    fetched = multi.get_presentation_slide_job(project(source), presentation["id"], job["id"])

    assert job["status"] == "succeeded"
    assert fetched["id"] == job["id"]
    assert fetched["proposal"]["patch"]["title"] == "理论基础"
    assert "证据边界" in fetched["proposal"]["patch"]["claim"]


def test_replace_rich_text_markdown_can_install_explicit_release_version(service):
    multi, source = service
    thesis = multi.list_documents(project(source))["documents"][0]
    replacement = (
        "# 文档说明\n\nv18正文。\n\n"
        "# 第一章 测试\n\n## 1.1 新结构\n\n正文内容。\n\n"
        "# 参考文献\n\n待补充。\n"
    )

    updated = multi.replace_rich_text_markdown(
        project(source),
        thesis["id"],
        replacement,
        "v18-release-sync",
        target_version=18,
    )

    workspace = multi.rich_workspace(project(source), thesis["id"])
    assert updated["revision"] == 18
    assert updated["source_checksum"] == hashlib.sha256(
        replacement.encode("utf-8")
    ).hexdigest()
    assert workspace["manifest"]["version"] == 18
    assert workspace["current_structure"]["version"] == "v18"
    assert workspace["current_structure"]["heading_count"] == 4
    assert source.read_text(encoding="utf-8") == replacement
    assert any(
        row["name"].startswith("v0001-before-v18-release-sync")
        for row in multi.versions(project(source), thesis["id"])
    )


def test_replace_rich_text_markdown_rejects_non_increasing_release_version(service):
    multi, source = service
    thesis = multi.list_documents(project(source))["documents"][0]

    with pytest.raises(Exception, match="必须高于当前版本"):
        multi.replace_rich_text_markdown(
            project(source),
            thesis["id"],
            "# 第一章 测试\n",
            "invalid-release-sync",
            target_version=1,
        )
